import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    DARK_GREEN = RGBColor(0x17, 0x6B, 0x4E)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    BLACK = RGBColor(0x00, 0x00, 0x00)
    DARK_TEXT = RGBColor(0x22, 0x22, 0x22)
    GRAY_BORDER = RGBColor(0xC0, 0xC5, 0xC0)

    screenshots_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), 'screenshots'))

    def add_header(slide, title_text):
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.65))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Cambria"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = DARK_GREEN

    def add_bullets(tf, points, font_size=14, space_after=4, line_spacing=1.15):
        for i, pt_text in enumerate(points):
            p = tf.add_paragraph() if i > 0 or len(tf.paragraphs[0].text) > 0 else tf.paragraphs[0]
            p.space_after = Pt(space_after)
            p.line_spacing = line_spacing
            
            if ":" in pt_text:
                parts = pt_text.split(":", 1)
                r1 = p.add_run()
                r1.text = "• " + parts[0] + ":"
                r1.font.name = "Calibri"
                r1.font.size = Pt(font_size)
                r1.font.bold = True
                r1.font.color.rgb = DARK_TEXT
                
                r2 = p.add_run()
                r2.text = parts[1]
                r2.font.name = "Calibri"
                r2.font.size = Pt(font_size)
                r2.font.bold = False
                r2.font.color.rgb = DARK_TEXT
            else:
                r = p.add_run()
                r.text = "• " + pt_text
                r.font.name = "Calibri"
                r.font.size = Pt(font_size)
                r.font.color.rgb = DARK_TEXT

    def add_picture_with_border(slide, filename, left, top, width=None, height=None):
        img_path = os.path.join(screenshots_dir, filename)
        if not os.path.exists(img_path):
            print(f"Warning: {filename} not found!")
            return None
        pic = slide.shapes.add_picture(img_path, left, top, width=width, height=height)
        pic.line.color.rgb = GRAY_BORDER
        pic.line.width = Pt(1)
        return pic

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_GREEN
    bg1.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "Skill Exchange"
    p.font.name = "Cambria"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(8)

    p = tf1.add_paragraph()
    p.text = "A platform connecting mentors and learners"
    p.font.name = "Calibri"
    p.font.size = Pt(22)
    p.font.color.rgb = RGBColor(0xDF, 0xF0, 0xE8)
    p.space_after = Pt(28)

    p = tf1.add_paragraph()
    p.text = "MERN Stack Project | B.Tech CSE, GLA University, Mathura | Internship at Codroid Hub Pvt Ltd"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.color.rgb = WHITE
    p.space_after = Pt(36)

    p = tf1.add_paragraph()
    p.text = "Team Members: __________________________________________________"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.color.rgb = WHITE
    p.space_after = Pt(12)

    p = tf1.add_paragraph()
    p.text = "Guide / Mentor: __________________________________________________"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 2: Problem and Objective
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Problem & Objective")

    tb2_left = slide2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.5), Inches(5.3))
    tf2_l = tb2_left.text_frame
    tf2_l.word_wrap = True
    p = tf2_l.paragraphs[0]
    p.text = "Problem Statement"
    p.font.name = "Cambria"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.space_after = Pt(14)
    add_bullets(tf2_l, [
        "Expensive Coaching: Commercial 1-on-1 tutoring platforms charge high hourly rates that college students cannot afford.",
        "Fragmented Channels: Students currently arrange peer help across WhatsApp, Discord, and Google Calendar without proper coordination.",
        "Ghosting & No-Shows: Informal peer sessions often fail due to lack of scheduling commitment and zero accountability.",
        "No Secure Escrow: Free swapping lacks reliability, while informal payments lack safety and dispute resolution."
    ], font_size=14, space_after=12, line_spacing=1.2)

    tb2_right = slide2.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
    tf2_r = tb2_right.text_frame
    tf2_r.word_wrap = True
    p = tf2_r.paragraphs[0]
    p.text = "Project Objectives"
    p.font.name = "Cambria"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.space_after = Pt(14)
    add_bullets(tf2_r, [
        "Centralized Peer Platform: Build a single web app connecting campus learners and mentors for verified 1-on-1 sessions.",
        "Flexible Scheduling: Provide date chips, duration picker, and slot negotiation with automated expiry windows.",
        "Real-Time Messaging: Integrate WebSocket-based chat for instant communication before and during sessions.",
        "Transparent Governance: Give administrators oversight over community members, session disputes, and payout ledgers."
    ], font_size=14, space_after=12, line_spacing=1.2)

    # ==========================================
    # SLIDE 3: Tech Stack
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Technology Stack")

    tech_items = [
        ("Frontend", "React 19 & Vite", [
            "Modern component-driven UI with Vite bundling for fast loading",
            "Client-side routing with React Router",
            "Lucide React iconography & custom clean CSS themes"
        ]),
        ("Backend", "Node.js & Express.js", [
            "RESTful API architecture handling auth, bookings, and profiles",
            "Modular controllers, routes, and middleware pipeline",
            "Input validation and CORS configuration"
        ]),
        ("Database", "MongoDB Atlas", [
            "NoSQL document database hosted on cloud Atlas",
            "Mongoose ODM for robust schema validation and indexing",
            "Stores users, bookings, messages, reviews, and logs"
        ]),
        ("Real-Time", "Socket.io", [
            "Bi-directional WebSocket connection for instant 1-on-1 chat",
            "Unread message tracking and active session gating",
            "Real-time presence and connection resilience"
        ]),
        ("Auth & Security", "JWT Authentication", [
            "Stateless token-based auth with accessToken stored securely",
            "Role-based authorization for Learner, Mentor, and Admin",
            "Bcrypt password hashing with salt rounds"
        ]),
        ("Integrations", "Razorpay & Google Meet", [
            "Escrow payment integration with Razorpay order API",
            "18% GST calculation and 3% platform commission logic",
            "Automatic Google Meet URL provisioning for confirmed calls"
        ])
    ]

    col_w = Inches(3.7)
    col_h = Inches(2.4)
    xs = [Inches(0.8), Inches(4.8), Inches(8.8)]
    ys = [Inches(1.5), Inches(4.2)]

    for idx, (cat, tech_name, points) in enumerate(tech_items):
        col_idx = idx % 3
        row_idx = idx // 3
        x = xs[col_idx]
        y = ys[row_idx]

        box = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, col_w, col_h)
        box.fill.background()
        box.line.color.rgb = GRAY_BORDER
        box.line.width = Pt(1)

        tb = slide3.shapes.add_textbox(x + Inches(0.15), y + Inches(0.12), col_w - Inches(0.3), col_h - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = f"{cat}: {tech_name}"
        p.font.name = "Cambria"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = DARK_GREEN
        p.space_after = Pt(6)

        for pt in points:
            p2 = tf.add_paragraph()
            p2.text = "• " + pt
            p2.font.name = "Calibri"
            p2.font.size = Pt(12)
            p2.font.color.rgb = DARK_TEXT
            p2.space_after = Pt(3)

    # ==========================================
    # SLIDE 4: System Architecture Diagram
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "System Architecture")

    # Box 1: React Client
    box1 = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.2), Inches(3.2), Inches(3.6))
    box1.fill.background()
    box1.line.color.rgb = BLACK
    box1.line.width = Pt(2.5)

    tf_b1 = box1.text_frame
    tf_b1.word_wrap = True
    p = tf_b1.paragraphs[0]
    p.text = "React Client (Frontend)\n"
    p.font.name = "Cambria"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.alignment = PP_ALIGN.CENTER

    for sp in ["Vite + React 19 UI", "Role Dashboards", "Booking Calendar", "Socket.io Client", "Local Storage Token"]:
        p2 = tf_b1.add_paragraph()
        p2.text = "• " + sp
        p2.font.name = "Calibri"
        p2.font.size = Pt(13)
        p2.font.color.rgb = BLACK

    # Arrow 1: REST API
    arr1 = slide4.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.3), Inches(2.8), Inches(1.3), Inches(0.45))
    arr1.fill.solid()
    arr1.fill.fore_color.rgb = BLACK
    arr1.line.color.rgb = BLACK

    tb_a1 = slide4.shapes.add_textbox(Inches(4.1), Inches(2.3), Inches(1.7), Inches(0.45))
    p = tb_a1.text_frame.paragraphs[0]
    p.text = "REST / HTTP"
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # Arrow 2: WebSockets
    arr2 = slide4.shapes.add_shape(MSO_SHAPE.LEFT_RIGHT_ARROW, Inches(4.3), Inches(4.6), Inches(1.3), Inches(0.45))
    arr2.fill.solid()
    arr2.fill.fore_color.rgb = BLACK
    arr2.line.color.rgb = BLACK

    tb_a2 = slide4.shapes.add_textbox(Inches(4.1), Inches(5.1), Inches(1.7), Inches(0.45))
    p = tb_a2.text_frame.paragraphs[0]
    p.text = "WebSocket (Chat)"
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # Box 2: Node.js / Express Server
    box2 = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.8), Inches(2.2), Inches(3.3), Inches(3.6))
    box2.fill.background()
    box2.line.color.rgb = BLACK
    box2.line.width = Pt(2.5)

    tf_b2 = box2.text_frame
    tf_b2.word_wrap = True
    p = tf_b2.paragraphs[0]
    p.text = "Express / Node.js API\n"
    p.font.name = "Cambria"
    p.font.size = Pt(18)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    for sp in ["JWT Auth Middleware", "Booking Lifecycle Engine", "Socket.io Chat Server", "Escrow & Fees Calculator", "Admin Governance API"]:
        p2 = tf_b2.add_paragraph()
        p2.text = "• " + sp
        p2.font.name = "Calibri"
        p2.font.size = Pt(13)
        p2.font.color.rgb = BLACK

    # Arrow 3: Mongoose ODM
    arr3 = slide4.shapes.add_shape(MSO_SHAPE.LEFT_RIGHT_ARROW, Inches(9.3), Inches(3.6), Inches(0.9), Inches(0.45))
    arr3.fill.solid()
    arr3.fill.fore_color.rgb = BLACK
    arr3.line.color.rgb = BLACK

    tb_a3 = slide4.shapes.add_textbox(Inches(9.1), Inches(3.05), Inches(1.3), Inches(0.45))
    p = tb_a3.text_frame.paragraphs[0]
    p.text = "Mongoose ODM"
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # Box 3: MongoDB Atlas
    box3 = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.4), Inches(2.2), Inches(2.3), Inches(3.6))
    box3.fill.background()
    box3.line.color.rgb = BLACK
    box3.line.width = Pt(2.5)

    tf_b3 = box3.text_frame
    tf_b3.word_wrap = True
    p = tf_b3.paragraphs[0]
    p.text = "MongoDB Atlas\n"
    p.font.name = "Cambria"
    p.font.size = Pt(18)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    for sp in ["Users & Profiles", "Bookings & Slots", "Chat Messages", "Reviews & Ratings", "Activity Logs"]:
        p2 = tf_b3.add_paragraph()
        p2.text = "• " + sp
        p2.font.name = "Calibri"
        p2.font.size = Pt(13)
        p2.font.color.rgb = BLACK

    # ==========================================
    # SLIDE 5: Home & Landing Page
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Home & Landing Page")

    tb5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2))
    tf5 = tb5.text_frame
    tf5.word_wrap = True
    add_bullets(tf5, [
        "Modern First Impression: Introduces students to the peer-to-peer knowledge sharing model with clear call-to-actions.",
        "Clear Exchange Proposition: Emphasizes two primary paths — 'Share what you know' as a mentor and 'Learn something new' as a student.",
        "Interactive Navigation: Fast access to Sign in, Register free, Theme mode toggle, and the experimental AI Matcher modal."
    ], font_size=15, space_after=14, line_spacing=1.2)

    add_picture_with_border(slide5, "01_home.png", Inches(5.5), Inches(1.6), width=Inches(7.0))

    # ==========================================
    # SLIDE 6: Authentication (Signup & Login)
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Authentication (Signup & Login)")

    tb6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.7), Inches(1.3))
    tf6 = tb6.text_frame
    tf6.word_wrap = True
    add_bullets(tf6, [
        "Role Selection: New users can easily choose between 'Learn Skills' and 'Teach & Mentor' during registration.",
        "Secure JWT Flow: Passwords are encrypted with bcrypt, and verified logins return JWT tokens saved in local storage.",
        "Direct Redirection: Directs authenticated users immediately to their respective learner or mentor dashboard."
    ], font_size=14, space_after=4, line_spacing=1.15)

    add_picture_with_border(slide6, "02a_login.png", Inches(0.8), Inches(2.65), width=Inches(5.6))
    add_picture_with_border(slide6, "02b_signup.png", Inches(6.9), Inches(2.65), width=Inches(5.6))

    # ==========================================
    # SLIDE 7: User Profile Management
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Profile Management (My Profile & Mentor Profile)")

    tb7 = slide7.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.7), Inches(1.3))
    tf7 = tb7.text_frame
    tf7.word_wrap = True
    add_bullets(tf7, [
        "Comprehensive Profile: Allows users to showcase bio, skills to teach, skills to learn, and available weekly time slots.",
        "Public Mentor Page: Displays verified mentor badges, hourly consultation rates, skill tags, and rating averages.",
        "Direct Action Triggers: Other students can start a 1-on-1 chat or jump directly into the session booking calendar."
    ], font_size=14, space_after=4, line_spacing=1.15)

    add_picture_with_border(slide7, "03a_my_profile.png", Inches(0.8), Inches(2.65), width=Inches(5.6))
    add_picture_with_border(slide7, "03b_other_user_profile.png", Inches(6.9), Inches(2.65), width=Inches(5.6))

    # ==========================================
    # SLIDE 8: Discover Mentors & Peers
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Discover Mentors & Peers")

    tb8 = slide8.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2))
    tf8 = tb8.text_frame
    tf8.word_wrap = True
    add_bullets(tf8, [
        "Searchable Directory: Lists verified mentors across various domains including React, Node.js, Python, and System Design.",
        "Skill & Rate Transparency: Each mentor card highlights key teaching skills, hourly fees, and average community star ratings.",
        "Fast Filtering: Instant live search lets students filter peers by name, specific technology tags, or availability."
    ], font_size=15, space_after=14, line_spacing=1.2)

    add_picture_with_border(slide8, "04_discover.png", Inches(5.5), Inches(1.6), width=Inches(7.0))

    # ==========================================
    # SLIDE 9: Learner Dashboard
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "Learner Dashboard")

    tb9 = slide9.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2))
    tf9 = tb9.text_frame
    tf9.word_wrap = True
    add_bullets(tf9, [
        "Learning Growth Workspace: Dedicated dashboard helping students keep track of their learning momentum and progress.",
        "Key Progress Metrics: Displays sessions attended, hours of live learning, mentors connected, and skills explored.",
        "Active Targets & Milestones: Visual tracking of topics like Node.js and System Design with milestone percentages."
    ], font_size=15, space_after=14, line_spacing=1.2)

    add_picture_with_border(slide9, "05_learner_dashboard.png", Inches(5.5), Inches(1.6), width=Inches(7.0))

    # ==========================================
    # SLIDE 10: Mentor Dashboard & Earnings
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "Mentor Dashboard & Earnings")

    tb10 = slide10.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.7), Inches(1.3))
    tf10 = tb10.text_frame
    tf10.word_wrap = True
    add_bullets(tf10, [
        "Teaching Studio: Live overview of teaching practice, learners taught, completed sessions, and skill level depth.",
        "Earnings & Settlement: Detailed ledger tracking lifetime revenue, settled payments, in-escrow funds, and pending payouts.",
        "UPI Payout Setup: Mentors can update and save their personal UPI ID directly for fast administrative settlements."
    ], font_size=14, space_after=4, line_spacing=1.15)

    add_picture_with_border(slide10, "06_mentor_dashboard.png", Inches(0.8), Inches(2.65), width=Inches(5.6))
    add_picture_with_border(slide10, "06b_mentor_earnings.png", Inches(6.9), Inches(2.65), width=Inches(5.6))

    # ==========================================
    # SLIDE 11: Session Booking Flow
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "Session Booking Flow")

    tb11 = slide11.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2))
    tf11 = tb11.text_frame
    tf11.word_wrap = True
    add_bullets(tf11, [
        "3-Step Scheduling: Clean wizard to pick session duration (30m, 45m, 1h, 1.5h, 2h) and an upcoming open calendar date.",
        "Morning/Evening Slots: Offers convenient time chips or custom time entry based on mentor availability.",
        "100% Escrow Protection: Clear fee calculation breakdown including base hourly rate, 18% GST, and secure escrow guarantee."
    ], font_size=15, space_after=14, line_spacing=1.2)

    add_picture_with_border(slide11, "07_booking_flow.png", Inches(5.5), Inches(1.6), width=Inches(7.0))

    # ==========================================
    # SLIDE 12: Real-Time Chat Messaging
    # ==========================================
    slide12 = prs.slides.add_slide(blank_layout)
    add_header(slide12, "Real-Time Chat Messaging")

    tb12 = slide12.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2))
    tf12 = tb12.text_frame
    tf12.word_wrap = True
    add_bullets(tf12, [
        "Socket.io Powered: Direct bi-directional messaging with instant delivery and live connectivity indicators.",
        "Conversation Management: Active chat list showing contact name, last message time, and unread counts.",
        "Spam Prevention: Enforces a 5-free direct message limit to prevent spam until a 1-on-1 mentorship session is scheduled."
    ], font_size=15, space_after=14, line_spacing=1.2)

    add_picture_with_border(slide12, "08_chat.png", Inches(5.5), Inches(1.6), width=Inches(7.0))

    # ==========================================
    # SLIDE 13: Ratings & Reviews System
    # ==========================================
    slide13 = prs.slides.add_slide(blank_layout)
    add_header(slide13, "Ratings & Reviews System")

    tb13 = slide13.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2))
    tf13 = tb13.text_frame
    tf13.word_wrap = True
    add_bullets(tf13, [
        "Verified Community Feedback: Learners can leave honest 1-5 star ratings and written reviews after completed sessions.",
        "Average Score Calculation: Real-time calculation of overall rating averages and review counts displayed on mentor cards.",
        "Trust & Quality Assurance: Motivates mentors to provide high-quality sessions and helps new learners pick trusted guides."
    ], font_size=15, space_after=14, line_spacing=1.2)

    add_picture_with_border(slide13, "09_ratings_reviews.png", Inches(5.5), Inches(1.6), width=Inches(7.0))

    # ==========================================
    # SLIDE 14: Admin Governance Dashboard
    # ==========================================
    slide14 = prs.slides.add_slide(blank_layout)
    add_header(slide14, "Admin Governance Dashboard")

    tb14 = slide14.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.7), Inches(1.3))
    tf14 = tb14.text_frame
    tf14.word_wrap = True
    add_bullets(tf14, [
        "Platform Overview KPIs: Live counts of total registered users, verified mentors, learners, and total completed sessions.",
        "Community Member Controls: Table allowing administrators to verify new accounts, suspend bad actors, or remove profiles.",
        "Financial & Session Audit: Tracks awaiting payments, escrow settlements, and generates full CSV ledger reports for audits."
    ], font_size=14, space_after=4, line_spacing=1.15)

    add_picture_with_border(slide14, "10_admin_dashboard.png", Inches(0.8), Inches(2.65), width=Inches(5.6))
    add_picture_with_border(slide14, "10b_admin_members.png", Inches(6.9), Inches(2.65), width=Inches(5.6))

    # ==========================================
    # SLIDE 15: Dark Theme & Responsive Design
    # ==========================================
    slide15 = prs.slides.add_slide(blank_layout)
    add_header(slide15, "Dark Theme & Responsive Design")

    tb15 = slide15.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.7), Inches(1.3))
    tf15 = tb15.text_frame
    tf15.word_wrap = True
    add_bullets(tf15, [
        "Built-in Dark Mode: User-friendly theme toggle switching the interface to an elegant dark emerald palette for night use.",
        "Mobile-First Experience: Fully responsive layout tested at 390x844 mobile resolution with compact navigation and cards.",
        "Adaptive Typography & Touch: Buttons, chips, and cards automatically adapt for smooth interaction across phones and laptops."
    ], font_size=14, space_after=4, line_spacing=1.15)

    add_picture_with_border(slide15, "11a_dark_home.png", Inches(0.8), Inches(2.65), width=Inches(6.8))
    add_picture_with_border(slide15, "12a_mobile_home.png", Inches(8.3), Inches(2.65), height=Inches(4.25))

    # ==========================================
    # SLIDE 16: Conclusion & Future Scope
    # ==========================================
    slide16 = prs.slides.add_slide(blank_layout)
    add_header(slide16, "Conclusion & Future Scope")

    tb16_l = slide16.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.5), Inches(5.3))
    tf16_l = tb16_l.text_frame
    tf16_l.word_wrap = True
    p = tf16_l.paragraphs[0]
    p.text = "Conclusion"
    p.font.name = "Cambria"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.space_after = Pt(14)
    add_bullets(tf16_l, [
        "Working MERN Solution: Successfully implemented a full-stack platform solving peer tutoring and knowledge exchange.",
        "Real-Time Collaboration: Integrated seamless Socket.io chat and calendar scheduling to eliminate coordination friction.",
        "Production Tested: Deployed live on Vercel and Render with live MongoDB Atlas database and role-based permissions.",
        "Undergraduate Team Work: Built cooperatively during our internship at Codroid Hub Pvt Ltd by GLA University students."
    ], font_size=14, space_after=12, line_spacing=1.2)

    tb16_r = slide16.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
    tf16_r = tb16_r.text_frame
    tf16_r.word_wrap = True
    p = tf16_r.paragraphs[0]
    p.text = "Future Scope"
    p.font.name = "Cambria"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.space_after = Pt(14)
    add_bullets(tf16_r, [
        "In-Browser Video Calling: Native WebRTC peer video calling directly inside the website without external Google Meet links.",
        "Automated Payment Webhooks: Live Razorpay automated webhooks for instant escrow fund releases and automated refunds.",
        "Push & Email Notifications: Instant browser push alerts and email reminders 15 minutes before scheduled sessions.",
        "Mobile App: Build native cross-platform iOS and Android apps using React Native or Flutter for quick on-the-go chat."
    ], font_size=14, space_after=12, line_spacing=1.2)

    # ==========================================
    # SLIDE 17: Thank You Slide
    # ==========================================
    slide17 = prs.slides.add_slide(blank_layout)
    bg17 = slide17.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg17.fill.solid()
    bg17.fill.fore_color.rgb = DARK_GREEN
    bg17.line.fill.background()

    tb17 = slide17.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(4.5))
    tf17 = tb17.text_frame
    tf17.word_wrap = True

    p = tf17.paragraphs[0]
    p.text = "Thank You!"
    p.font.name = "Cambria"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(12)

    p = tf17.add_paragraph()
    p.text = "Questions & Discussion"
    p.font.name = "Calibri"
    p.font.size = Pt(22)
    p.font.color.rgb = RGBColor(0xDF, 0xF0, 0xE8)
    p.space_after = Pt(36)

    p = tf17.add_paragraph()
    p.text = "Project Repository: github.com/CosmicCoder06/Skill-Exchange"
    p.font.name = "Calibri"
    p.font.size = Pt(18)
    p.font.color.rgb = WHITE
    p.space_after = Pt(12)

    p = tf17.add_paragraph()
    p.text = "Live Demo: https://skill-exchange-one-eta.vercel.app/"
    p.font.name = "Calibri"
    p.font.size = Pt(18)
    p.font.color.rgb = WHITE

    output_path = os.path.abspath(os.path.join(os.path.dirname(__file__), 'Skill_Exchange_Presentation.pptx'))
    prs.save(output_path)
    print(f"Presentation saved successfully at: {output_path}")

if __name__ == "__main__":
    create_presentation()
