import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_FILE = os.path.abspath("Skill_Exchange_College_Submission_Presentation.pptx")
SCREENSHOTS_DIR = os.path.abspath("screenshots")
EXPORT_DIR = os.path.abspath("submission_slide_exports")

# Brand Color Palette
DARK_GREEN = RGBColor(0x17, 0x6B, 0x4E)     # Primary Forest Green
DEEP_GREEN = RGBColor(0x0F, 0x47, 0x33)     # Deep Forest Accent
LIGHT_BG = RGBColor(0xF8, 0xFB, 0xF9)       # Clean Page Background
CARD_BG = RGBColor(0xFF, 0xFF, 0xFF)        # White Card Fill
CARD_ACCENT = RGBColor(0xF2, 0xF7, 0xF4)    # Mint Accent Box
TEXT_DARK = RGBColor(0x1A, 0x20, 0x2C)      # Primary Heading Dark Slate
TEXT_BODY = RGBColor(0x2D, 0x37, 0x48)      # Readable Paragraph Text
TEXT_MUTED = RGBColor(0x5A, 0x65, 0x78)     # Secondary Subtext
BORDER_COLOR = RGBColor(0xD1, 0xE3, 0xD9)   # Mint Outline Border
BORDER_LIGHT = RGBColor(0xE2, 0xE8, 0xF0)   # Subtle Card Border
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TAG_GREEN_BG = RGBColor(0xE6, 0xF4, 0xEA)   # Badge Fill

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]
TOTAL_SLIDES = 10

def set_slide_bg(slide, color=LIGHT_BG):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, title_text, category_tag="SKILL EXCHANGE PLATFORM"):
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(11.7), Inches(0.28))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category_tag.upper()
    p_tag.font.name = "Calibri"
    p_tag.font.size = Pt(9.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = DARK_GREEN

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.66), Inches(11.7), Inches(0.55))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Cambria"
    p_title.font.size = Pt(25)
    p_title.font.bold = True
    p_title.font.color.rgb = DARK_GREEN

def add_footer(slide, slide_num):
    # Left: Project & College Tagline
    tb_l = slide.shapes.add_textbox(Inches(0.8), Inches(7.08), Inches(8.0), Inches(0.28))
    tf_l = tb_l.text_frame
    tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = "Learn together.  Share knowledge.  Grow together.  •  Team CosmicCoders (Group 13)"
    p_l.font.name = "Calibri"
    p_l.font.size = Pt(9.5)
    p_l.font.color.rgb = TEXT_MUTED

    # Right: Slide Number
    tb_r = slide.shapes.add_textbox(Inches(10.5), Inches(7.08), Inches(2.033), Inches(0.28))
    tf_r = tb_r.text_frame
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    p_r = tf_r.paragraphs[0]
    p_r.text = f"{slide_num:02d} / {TOTAL_SLIDES:02d}"
    p_r.font.name = "Calibri"
    p_r.font.size = Pt(10)
    p_r.font.bold = True
    p_r.font.color.rgb = DARK_GREEN
    p_r.alignment = PP_ALIGN.RIGHT

def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card

def add_image_bordered(slide, filename, left, top, width, height):
    # Check in screenshots dir first, then docs/screenshots dir
    path = os.path.join(SCREENSHOTS_DIR, filename)
    if not os.path.exists(path):
        alt_path = os.path.join(os.path.abspath("docs/screenshots"), filename)
        if os.path.exists(alt_path):
            path = alt_path
        else:
            print(f"Warning: Screenshot {filename} missing!")
            return None
    pic = slide.shapes.add_picture(path, left, top, width=width, height=height)
    pic.line.color.rgb = BORDER_COLOR
    pic.line.width = Pt(1)
    return pic

def add_speaker_note(slide, note_text):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = note_text

# =================================================================================================
# SLIDE 1: Title Page (Exact UI replication as user's reference)
# =================================================================================================
def build_slide_1():
    s = prs.slides.add_slide(blank_layout)
    set_slide_bg(s, LIGHT_BG)

    # Top Bar: TEAM COSMICCODERS | GROUP 13
    top_box = s.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(6.0), Inches(0.3))
    tf_top = top_box.text_frame
    tf_top.word_wrap = True
    tf_top.margin_left = tf_top.margin_top = tf_top.margin_right = tf_top.margin_bottom = 0
    p_top = tf_top.paragraphs[0]
    p_top.text = "TEAM COSMICCODERS  |  GROUP 13"
    p_top.font.name = "Calibri"
    p_top.font.size = Pt(11)
    p_top.font.bold = True
    p_top.font.color.rgb = DARK_GREEN

    # Logo & Title
    logo_path = os.path.join(SCREENSHOTS_DIR, "logo.png")
    if os.path.exists(logo_path):
        s.shapes.add_picture(logo_path, Inches(0.8), Inches(1.2), Inches(0.9), Inches(0.9))

    title_box = s.shapes.add_textbox(Inches(1.8), Inches(1.15), Inches(5.0), Inches(1.0))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = "SkillExchange"
    p_t.font.name = "Cambria"
    p_t.font.size = Pt(44)
    p_t.font.bold = True
    p_t.font.color.rgb = DARK_GREEN

    # Subtitle & Badges
    desc_box = s.shapes.add_textbox(Inches(0.8), Inches(2.25), Inches(6.0), Inches(2.0))
    tf_d = desc_box.text_frame
    tf_d.word_wrap = True
    tf_d.margin_left = tf_d.margin_top = tf_d.margin_right = tf_d.margin_bottom = 0
    
    p_sub = tf_d.paragraphs[0]
    p_sub.text = "A platform connecting mentors and learners"
    p_sub.font.name = "Calibri"
    p_sub.font.size = Pt(18)
    p_sub.font.bold = True
    p_sub.font.color.rgb = TEXT_DARK
    p_sub.space_after = Pt(12)

    p_tag = tf_d.add_paragraph()
    p_tag.text = "MERN STACK PROJECT  •  B.TECH CSE"
    p_tag.font.name = "Calibri"
    p_tag.font.size = Pt(13)
    p_tag.font.bold = True
    p_tag.font.color.rgb = DARK_GREEN
    p_tag.space_after = Pt(10)

    p_col = tf_d.add_paragraph()
    p_col.text = "GLA University, Mathura  |  Internship at Codroid Hub Pvt. Ltd."
    p_col.font.name = "Calibri"
    p_col.font.size = Pt(13)
    p_col.font.color.rgb = TEXT_BODY
    p_col.space_after = Pt(20)

    p_desc = tf_d.add_paragraph()
    p_desc.text = "A centralized, full-stack peer mentorship web application built with React 19, Node.js, Express, MongoDB Atlas, and Socket.io for frictionless peer tutoring and verifiable skill exchange."
    p_desc.font.name = "Calibri"
    p_desc.font.size = Pt(11)
    p_desc.font.color.rgb = TEXT_MUTED

    # Right Column: Clean White Card for Team Details
    card_w = Inches(5.6)
    card_h = Inches(5.2)
    card_x = Inches(6.9)
    card_y = Inches(1.1)
    add_card(s, card_x, card_y, card_w, card_h, bg_color=CARD_BG, border_color=BORDER_COLOR)

    tb_card = s.shapes.add_textbox(card_x + Inches(0.4), card_y + Inches(0.35), card_w - Inches(0.8), card_h - Inches(0.7))
    tf_c = tb_card.text_frame
    tf_c.word_wrap = True

    p0 = tf_c.paragraphs[0]
    p0.text = "TEAM COSMICCODERS  •  GROUP 13"
    p0.font.name = "Calibri"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = DARK_GREEN
    p0.space_after = Pt(14)

    p1 = tf_c.add_paragraph()
    p1.text = "Team Leader:"
    p1.font.name = "Calibri"
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_DARK

    p1_line = tf_c.add_paragraph()
    p1_line.text = "________________________________________"
    p1_line.font.name = "Calibri"
    p1_line.font.size = Pt(11)
    p1_line.font.color.rgb = TEXT_MUTED
    p1_line.space_after = Pt(12)

    p2 = tf_c.add_paragraph()
    p2.text = "Team Members:"
    p2.font.name = "Calibri"
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_DARK

    members = [
        "1. ______________________",
        "2. ______________________",
        "3. ______________________",
        "4. ______________________"
    ]
    for m in members:
        pm = tf_c.add_paragraph()
        pm.text = m
        pm.font.name = "Calibri"
        pm.font.size = Pt(10.5)
        pm.font.color.rgb = TEXT_MUTED
        pm.space_after = Pt(3)

    p3 = tf_c.add_paragraph()
    p3.space_before = Pt(8)
    p3.text = "Project Guide / Mentor:"
    p3.font.name = "Calibri"
    p3.font.size = Pt(11)
    p3.font.bold = True
    p3.font.color.rgb = TEXT_DARK

    p3_line = tf_c.add_paragraph()
    p3_line.text = "________________________________________"
    p3_line.font.name = "Calibri"
    p3_line.font.size = Pt(11)
    p3_line.font.color.rgb = TEXT_MUTED

    # Bottom Tagline
    tb_bot = s.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(8.0), Inches(0.3))
    tf_b = tb_bot.text_frame
    p_bot = tf_b.paragraphs[0]
    p_bot.text = "Learn together.  Share knowledge.  Grow together."
    p_bot.font.name = "Calibri"
    p_bot.font.size = Pt(11)
    p_bot.font.italic = True
    p_bot.font.color.rgb = TEXT_MUTED

    tb_num = s.shapes.add_textbox(Inches(11.5), Inches(6.95), Inches(1.0), Inches(0.3))
    tf_num = tb_num.text_frame
    p_num = tf_num.paragraphs[0]
    p_num.text = "01"
    p_num.font.name = "Calibri"
    p_num.font.size = Pt(14)
    p_num.font.bold = True
    p_num.font.color.rgb = DARK_GREEN
    p_num.alignment = PP_ALIGN.RIGHT

    add_speaker_note(s,
        "Good morning respected professors and project evaluators. Today, Team CosmicCoders (Group 13) presents our college capstone project: "
        "'SkillExchange' — A centralized, full-stack peer mentorship and skill-sharing web platform developed using the modern MERN stack. "
        "This project bridges the gap between students seeking targeted academic/technical guidance and skilled senior peers willing to mentor."
    )

# =================================================================================================
# SLIDE 2: Problem Statement & Project Objectives
# =================================================================================================
def build_slide_2():
    s = prs.slides.add_slide(blank_layout)
    set_slide_bg(s)
    add_header(s, "Problem Statement & Project Objectives", "CONTEXT & MOTIVATION")
    add_footer(s, 2)

    col_w = Inches(5.7)
    col_h = Inches(5.4)

    # Left Column: Problem Statement Card
    add_card(s, Inches(0.8), Inches(1.45), col_w, col_h, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb_l = s.shapes.add_textbox(Inches(1.05), Inches(1.65), col_w - Inches(0.5), col_h - Inches(0.4))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p_lh = tf_l.paragraphs[0]
    p_lh.text = "The Campus Peer Learning Gap"
    p_lh.font.name = "Cambria"
    p_lh.font.size = Pt(16)
    p_lh.font.bold = True
    p_lh.font.color.rgb = DARK_GREEN
    p_lh.space_after = Pt(12)

    problems = [
        ("Expensive Commercial Tutoring", "Existing 1-on-1 coaching platforms (Preply, Chegg) charge commercial rates (₹500–₹2500/hr) unaffordable for college students."),
        ("Fragmented Informal Channels", "Students arrange peer help across WhatsApp, Discord, and Google Calendar without proper coordination or centralized history."),
        ("Ghosting & Zero Accountability", "Informal peer sessions suffer frequent no-shows, delayed cancellations, and lack of scheduling commitment."),
        ("Lack of Financial Transparency", "Students lack transparent pricing, receipts, and a secure way to track peer honorariums and mentor payouts.")
    ]
    for title, desc in problems:
        p = tf_l.add_paragraph()
        p.space_after = Pt(10)
        r0 = p.add_run()
        r0.text = "• " + title + ": "
        r0.font.name = "Calibri"
        r0.font.size = Pt(10.5)
        r0.font.bold = True
        r0.font.color.rgb = TEXT_DARK
        r1 = p.add_run()
        r1.text = desc
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        r1.font.color.rgb = TEXT_BODY

    # Right Column: Project Objectives Card
    add_card(s, Inches(6.83), Inches(1.45), col_w, col_h, bg_color=CARD_ACCENT, border_color=BORDER_COLOR)
    tb_r = s.shapes.add_textbox(Inches(7.08), Inches(1.65), col_w - Inches(0.5), col_h - Inches(0.4))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p_rh = tf_r.paragraphs[0]
    p_rh.text = "Project Objectives & Solutions"
    p_rh.font.name = "Cambria"
    p_rh.font.size = Pt(16)
    p_rh.font.bold = True
    p_rh.font.color.rgb = DARK_GREEN
    p_rh.space_after = Pt(12)

    objectives = [
        ("Unified Peer Platform", "Build a dedicated web portal where campus learners and mentors connect through verified university profiles."),
        ("Deterministic 5-Step Scheduling", "Provide intuitive duration chips, date selectors, slot validation, and a strict 4-hour automatic request expiration."),
        ("Real-Time WebSocket Communication", "Integrate Socket.io messaging for instant chat with a 5-free message anti-spam quota unlocked upon session booking."),
        ("Transparent Wallet & Receipts", "Implement dual-balance wallet architecture (Top-Up vs Earned), simulated escrow, and 0% platform fee receipts."),
        ("Institutional Governance", "Provide administrators with platform KPIs, member verification/suspension controls, and auditable activity logs.")
    ]
    for title, desc in objectives:
        p = tf_r.add_paragraph()
        p.space_after = Pt(8)
        r0 = p.add_run()
        r0.text = "✓ " + title + ": "
        r0.font.name = "Calibri"
        r0.font.size = Pt(10.5)
        r0.font.bold = True
        r0.font.color.rgb = DARK_GREEN
        r1 = p.add_run()
        r1.text = desc
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        r1.font.color.rgb = TEXT_BODY

    add_speaker_note(s,
        "Here we discuss the core problem statement. College students often struggle to find timely academic help. "
        "Commercial tutoring platforms are far too expensive, while informal WhatsApp arrangements lead to high ghosting rates. "
        "SkillExchange addresses this through a centralized MERN web application offering verified mentor profiles, deterministic scheduling, "
        "real-time messaging with anti-spam limits, and a transparent simulated escrow wallet."
    )

# =================================================================================================
# SLIDE 3: System Architecture & Modern Technology Stack
# =================================================================================================
def build_slide_3():
    s = prs.slides.add_slide(blank_layout)
    set_slide_bg(s)
    add_header(s, "System Architecture & MERN Technology Stack", "TECHNICAL FOUNDATION")
    add_footer(s, 3)

    # 4 Architecture Flow Cards across the top
    flow_cards = [
        ("1. Client Frontend", "React 19 & Vite", [
            "Component-driven SPA architecture",
            "Tailwind & custom CSS with Light/Dark mode",
            "Role-based routing & Socket.io client"
        ]),
        ("2. Application API", "Node.js & Express.js", [
            "RESTful JSON controllers & routes",
            "JWT bearer auth & RBAC middleware",
            "Booking lifecycle engine & fee calculators"
        ]),
        ("3. Real-Time Engine", "Socket.io WebSockets", [
            "Bi-directional event-driven messaging",
            "Room-based chat sessions & typing events",
            "Unread badge counters & online status"
        ]),
        ("4. Database Cloud", "MongoDB Atlas", [
            "Mongoose ODM for robust schema indexing",
            "Collections: Users, Bookings, Messages,",
            "Reviews, Wallets, and Audit Logs"
        ])
    ]

    card_w = Inches(2.78)
    card_h = Inches(2.35)
    for i, (title, sub, pts) in enumerate(flow_cards):
        cx = Inches(0.8) + i * Inches(2.97)
        cy = Inches(1.45)
        add_card(s, cx, cy, card_w, card_h, bg_color=CARD_BG, border_color=BORDER_COLOR)

        tb = s.shapes.add_textbox(cx + Inches(0.18), cy + Inches(0.15), card_w - Inches(0.36), card_h - Inches(0.30))
        tf = tb.text_frame
        tf.word_wrap = True

        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = "Cambria"
        pt.font.size = Pt(12)
        pt.font.bold = True
        pt.font.color.rgb = DARK_GREEN
        pt.space_after = Pt(2)

        psub = tf.add_paragraph()
        psub.text = sub
        psub.font.name = "Calibri"
        psub.font.size = Pt(9.5)
        psub.font.bold = True
        psub.font.color.rgb = TEXT_DARK
        psub.space_after = Pt(6)

        for pt_txt in pts:
            p = tf.add_paragraph()
            p.space_after = Pt(2)
            p.text = "• " + pt_txt
            p.font.name = "Calibri"
            p.font.size = Pt(8.5)
            p.font.color.rgb = TEXT_BODY

    # Bottom Half: 2 Detailed Cards (Architecture Protocol & Security)
    bot_w = Inches(5.75)
    bot_h = Inches(2.90)
    bot_y = Inches(3.95)

    # Left: Data Flow & Request Lifecycle
    add_card(s, Inches(0.8), bot_y, bot_w, bot_h, bg_color=CARD_ACCENT, border_color=BORDER_COLOR)
    tb_df = s.shapes.add_textbox(Inches(1.05), bot_y + Inches(0.15), bot_w - Inches(0.5), bot_h - Inches(0.3))
    tf_df = tb_df.text_frame
    tf_df.word_wrap = True
    
    pdf_h = tf_df.paragraphs[0]
    pdf_h.text = "End-to-End Data Flow Pipeline"
    pdf_h.font.name = "Cambria"
    pdf_h.font.size = Pt(13)
    pdf_h.font.bold = True
    pdf_h.font.color.rgb = DARK_GREEN
    pdf_h.space_after = Pt(6)

    flow_steps = [
        ("Client Request", "React client dispatches REST HTTP requests via Axios interceptor with Bearer JWT."),
        ("Auth Validation", "Express authMiddleware intercepts request, verifies JWT signature, and attaches req.user."),
        ("Business Logic", "Dedicated controllers (booking, user, payment) execute slot checking, state changes, or calculations."),
        ("Persistence Layer", "Mongoose models validate constraints, timestamps, and save documents into MongoDB Atlas.")
    ]
    for st, sd in flow_steps:
        p = tf_df.add_paragraph()
        p.space_after = Pt(3)
        r0 = p.add_run()
        r0.text = "→ " + st + ": "
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = TEXT_DARK
        r1 = p.add_run()
        r1.text = sd
        r1.font.name = "Calibri"
        r1.font.size = Pt(9)
        r1.font.color.rgb = TEXT_BODY

    # Right: Security & Integration Specs
    add_card(s, Inches(6.78), bot_y, bot_w, bot_h, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb_sc = s.shapes.add_textbox(Inches(7.03), bot_y + Inches(0.15), bot_w - Inches(0.5), bot_h - Inches(0.3))
    tf_sc = tb_sc.text_frame
    tf_sc.word_wrap = True

    psc_h = tf_sc.paragraphs[0]
    psc_h.text = "Security Standards & Integrations"
    psc_h.font.name = "Cambria"
    psc_h.font.size = Pt(13)
    psc_h.font.bold = True
    psc_h.font.color.rgb = DARK_GREEN
    psc_h.space_after = Pt(6)

    sec_items = [
        ("Cryptographic Hashing", "Passwords hashed using 10-round salted Bcrypt prior to database insertion."),
        ("Stateless JWT Authorization", "HS256 tokens carrying role claims (Learner, Mentor, Admin) preventing privilege escalation."),
        ("Payment Checkout Gateway", "Razorpay integration supporting test card/UPI checkout, combined with wallet balances."),
        ("Simulated Escrow Protection", "Transparent fee model: 0% platform fee, 0% tax, ₹0 extra convenience fees credited on completion.")
    ]
    for st, sd in sec_items:
        p = tf_sc.add_paragraph()
        p.space_after = Pt(3)
        r0 = p.add_run()
        r0.text = "🔒 " + st + ": "
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = TEXT_DARK
        r1 = p.add_run()
        r1.text = sd
        r1.font.name = "Calibri"
        r1.font.size = Pt(9)
        r1.font.color.rgb = TEXT_BODY

    add_speaker_note(s,
        "This slide illustrates our full-stack architecture. We selected the MERN stack for high developer productivity and performance. "
        "The frontend is built with React 19 and Vite. The backend utilizes Express.js and Node.js with strict JWT authentication middleware. "
        "Data is persisted in MongoDB Atlas via Mongoose ODM, and real-time messaging is powered by Socket.io."
    )

# =================================================================================================
# SLIDE 4: Module 01: Authentication, Profiles & Skill Management
# =================================================================================================
def build_slide_4():
    s = prs.slides.add_slide(blank_layout)
    set_slide_bg(s)
    add_header(s, "Module 01: Authentication, Profiles & Skill Management", "CORE FUNCTIONAL MODULE")
    add_footer(s, 4)

    left_w = Inches(5.60)
    top_y = Inches(1.45)

    # Left Column: Features Card
    card_h = Inches(5.40)
    add_card(s, Inches(0.8), top_y, left_w, card_h, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb = s.shapes.add_textbox(Inches(1.05), top_y + Inches(0.20), left_w - Inches(0.5), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    ph = tf.paragraphs[0]
    ph.text = "Identity, Access & Profile Capabilities"
    ph.font.name = "Cambria"
    ph.font.size = Pt(14)
    ph.font.bold = True
    ph.font.color.rgb = DARK_GREEN
    ph.space_after = Pt(8)

    features = [
        ("Dual-Role Registration", "Users choose between 'Learn Skills' and 'Teach & Mentor' during signup with specific validation schemas."),
        ("Cryptographic Security", "Passwords hashed with Bcrypt (10 salt rounds); validated sessions return signed JWT bearer tokens stored in localStorage."),
        ("Learner Profile & Goals", "Learners maintain academic bios, learning goals (React, Node.js, System Design), and track profile completion percentages."),
        ("Mentor Skill Depth & Rates", "Mentors showcase domain expertise, granular skill proficiency ratings (Beginner to Expert), and set hourly rates (₹/hr)."),
        ("Public Profile & Trust Badges", "Public profile showcase allows peers to inspect bios, verified university badges, average star ratings, and community reviews.")
    ]
    for hd, desc in features:
        p = tf.add_paragraph()
        p.space_after = Pt(6)
        r0 = p.add_run()
        r0.text = "• " + hd + ": "
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = TEXT_DARK
        r1 = p.add_run()
        r1.text = desc
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = TEXT_BODY

    # Technical Specs Mini Box inside left column
    p_tech = tf.add_paragraph()
    p_tech.space_before = Pt(8)
    p_tech.text = "APIs: POST /api/register, POST /api/login, GET/PUT /api/profile | Model: UserSchema"
    p_tech.font.name = "Calibri"
    p_tech.font.size = Pt(9)
    p_tech.font.bold = True
    p_tech.font.color.rgb = DARK_GREEN

    # Right Column: 2 Stacked Screenshots (Signup & Profile)
    img_left = Inches(6.65)
    img_w = Inches(5.88)
    img_h1 = Inches(2.62)
    img_h2 = Inches(2.62)

    add_image_bordered(s, "02b_signup.png", img_left, top_y, img_w, img_h1)
    add_image_bordered(s, "03a_my_profile.png", img_left, top_y + img_h1 + Inches(0.16), img_w, img_h2)

    add_speaker_note(s,
        "Module 1 covers User Onboarding, Authentication, and Profiles. We offer dual registration for learners and mentors. "
        "Authentication is secured via 10-round salted Bcrypt password encryption and JWT stateless tokens. "
        "Learners specify skills they wish to learn, while mentors configure their hourly rates and proficiency levels."
    )

# =================================================================================================
# SLIDE 5: Module 02: Peer Discovery & Intelligent AI Matcher
# =================================================================================================
def build_slide_5():
    s = prs.slides.add_slide(blank_layout)
    set_slide_bg(s)
    add_header(s, "Module 02: Peer Discovery & Intelligent AI Matcher", "CORE FUNCTIONAL MODULE")
    add_footer(s, 5)

    left_w = Inches(5.60)
    top_y = Inches(1.45)

    # Left Column: Features Card
    card_h = Inches(5.40)
    add_card(s, Inches(0.8), top_y, left_w, card_h, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb = s.shapes.add_textbox(Inches(1.05), top_y + Inches(0.20), left_w - Inches(0.5), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    ph = tf.paragraphs[0]
    ph.text = "Search, Filtering & AI Compatibility Engine"
    ph.font.name = "Cambria"
    ph.font.size = Pt(14)
    ph.font.bold = True
    ph.font.color.rgb = DARK_GREEN
    ph.space_after = Pt(8)

    features = [
        ("Searchable Community Catalog", "Fast indexed directory listing verified mentors across modern technical stacks (React, Node.js, Python, System Design)."),
        ("Live Multi-Parametric Filters", "Instant client-side filtering by mentor name, skill tags, availability days, and fee ranges without page reloads."),
        ("Card Transparency", "Each mentor card highlights teaching skills, verified university badge, hourly rate (e.g. ₹500/hr), and star reviews."),
        ("Gemini-Powered AI Matcher", "Dedicated 'AI Matcher' modal analyzes learner's desired skills, learning rhythm, and mentor schedule compatibility to recommend optimal peers."),
        ("One-Click Profile Inspection", "Learners can seamlessly navigate from discovery to detailed mentor profiles to review student testimonials and initiate booking.")
    ]
    for hd, desc in features:
        p = tf.add_paragraph()
        p.space_after = Pt(6)
        r0 = p.add_run()
        r0.text = "• " + hd + ": "
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = TEXT_DARK
        r1 = p.add_run()
        r1.text = desc
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = TEXT_BODY

    p_tech = tf.add_paragraph()
    p_tech.space_before = Pt(8)
    p_tech.text = "APIs: GET /api/users/mentors, POST /api/ai/match | Components: DiscoverPage.jsx, AIMatcherModal.jsx"
    p_tech.font.name = "Calibri"
    p_tech.font.size = Pt(9)
    p_tech.font.bold = True
    p_tech.font.color.rgb = DARK_GREEN

    # Right Column: 2 Stacked Screenshots (Discover & AI Matcher)
    img_left = Inches(6.65)
    img_w = Inches(5.88)
    img_h1 = Inches(2.62)
    img_h2 = Inches(2.62)

    add_image_bordered(s, "04_discover.png", img_left, top_y, img_w, img_h1)
    add_image_bordered(s, "04b_ai_matcher.png", img_left, top_y + img_h1 + Inches(0.16), img_w, img_h2)

    add_speaker_note(s,
        "Module 2 is Peer Discovery & the AI Matcher. Learners can browse mentors by technology tags or search by name. "
        "A key innovation is the Gemini-powered AI Matcher, which allows students to describe what they want to learn and automatically "
        "pairs them with the best suited mentor based on skill overlap, schedule compatibility, and community ratings."
    )

# =================================================================================================
# SLIDE 6: Module 03: 5-Stage Session Booking & Slot Negotiation
# =================================================================================================
def build_slide_6():
    s = prs.slides.add_slide(blank_layout)
    set_slide_bg(s)
    add_header(s, "Module 03: 5-Stage Session Booking & Slot Negotiation", "SIGNATURE FEATURE")
    add_footer(s, 6)

    left_w = Inches(5.60)
    top_y = Inches(1.45)

    # Left Column: Features Card
    card_h = Inches(5.40)
    add_card(s, Inches(0.8), top_y, left_w, card_h, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb = s.shapes.add_textbox(Inches(1.05), top_y + Inches(0.20), left_w - Inches(0.5), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    ph = tf.paragraphs[0]
    ph.text = "Deterministic 5-Step Scheduling Lifecycle"
    ph.font.name = "Cambria"
    ph.font.size = Pt(14)
    ph.font.bold = True
    ph.font.color.rgb = DARK_GREEN
    ph.space_after = Pt(8)

    steps = [
        ("Step 1: Duration & Date Selection", "Learner selects session duration (30m, 45m, 1h, 1.5h, 2h) and upcoming date chips with automatic cost calculation."),
        ("Step 2: Real-Time Slot Validation", "System checks mentor's active calendar availability to eliminate double-booking conflicts before submission."),
        ("Step 3: Mentor Response Window", "Mentor receives the request and can Approve, Propose Alternate Slots, or Decline with custom feedback."),
        ("Step 4: Automatic 4-Hour Expiration", "Unanswered pending requests automatically expire after 4 hours, preventing stale bookings and scheduling deadlock."),
        ("Step 5: Escrow Payment & Confirmation", "Payment is authorized via wallet or checkout and held securely in escrow until the session is successfully conducted.")
    ]
    for hd, desc in steps:
        p = tf.add_paragraph()
        p.space_after = Pt(6)
        r0 = p.add_run()
        r0.text = "• " + hd + ": "
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = TEXT_DARK
        r1 = p.add_run()
        r1.text = desc
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = TEXT_BODY

    p_tech = tf.add_paragraph()
    p_tech.space_before = Pt(8)
    p_tech.text = "APIs: POST /api/bookings, PUT /api/bookings/:id/status | Model: BookingSchema (Enum status)"
    p_tech.font.name = "Calibri"
    p_tech.font.size = Pt(9)
    p_tech.font.bold = True
    p_tech.font.color.rgb = DARK_GREEN

    # Right Column: 2 Stacked Screenshots (Booking Flow & Payment Checkout)
    img_left = Inches(6.65)
    img_w = Inches(5.88)
    img_h1 = Inches(2.62)
    img_h2 = Inches(2.62)

    add_image_bordered(s, "07_booking_flow.png", img_left, top_y, img_w, img_h1)
    add_image_bordered(s, "07_payment_checkout.png", img_left, top_y + img_h1 + Inches(0.16), img_w, img_h2)

    add_speaker_note(s,
        "Module 3 is our signature Session Booking Flow. Unlike informal messaging, our platform implements a 5-step deterministic lifecycle. "
        "The learner picks duration and date chips, the system validates mentor slots, and the mentor can approve or suggest counter-slots. "
        "Crucially, any request left unaddressed for 4 hours automatically expires to avoid stale scheduling. Funds are held in escrow."
    )

# =================================================================================================
# SLIDE 7: Module 04: Real-Time Chat & Wallet / Payment Engine
# =================================================================================================
def build_slide_7():
    s = prs.slides.add_slide(blank_layout)
    set_slide_bg(s)
    add_header(s, "Module 04: Real-Time Chat & Wallet / Payment Engine", "CORE FUNCTIONAL MODULE")
    add_footer(s, 7)

    left_w = Inches(5.60)
    top_y = Inches(1.45)

    # Left Column: Features Card
    card_h = Inches(5.40)
    add_card(s, Inches(0.8), top_y, left_w, card_h, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb = s.shapes.add_textbox(Inches(1.05), top_y + Inches(0.20), left_w - Inches(0.5), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    ph = tf.paragraphs[0]
    ph.text = "WebSocket Messaging & Dual-Balance Wallet"
    ph.font.name = "Cambria"
    ph.font.size = Pt(14)
    ph.font.bold = True
    ph.font.color.rgb = DARK_GREEN
    ph.space_after = Pt(8)

    features = [
        ("Socket.io Bi-Directional Messaging", "Sub-100ms instant message delivery with live connection indicators, typing events, and unread counters."),
        ("Anti-Spam 5-Message Quota", "To prevent mentor harassment, free direct chat is strictly capped at 5 messages until a 1-on-1 session is booked; booking unlocks unlimited messaging."),
        ("Dual-Balance Wallet Architecture", "Separates spendable 'Top-Up Balance' (for purchasing sessions) from 'Earned Balance' (mentor revenue credited upon session completion)."),
        ("Transparent Pricing & 0% Fees", "Clear session receipts with 0% platform deductions, 0% convenience fees, and 0% taxes; 100% of agreed fees credit to the mentor."),
        ("Manual Settlement & Payout Tracking", "Mentors configure registered UPI IDs to receive disbursed earnings tracked through life-time revenue and pending settlement cards.")
    ]
    for hd, desc in features:
        p = tf.add_paragraph()
        p.space_after = Pt(6)
        r0 = p.add_run()
        r0.text = "• " + hd + ": "
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = TEXT_DARK
        r1 = p.add_run()
        r1.text = desc
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = TEXT_BODY

    p_tech = tf.add_paragraph()
    p_tech.space_before = Pt(8)
    p_tech.text = "Sockets: send_message, receive_message, typing | APIs: GET /api/wallet, POST /api/wallet/topup"
    p_tech.font.name = "Calibri"
    p_tech.font.size = Pt(9)
    p_tech.font.bold = True
    p_tech.font.color.rgb = DARK_GREEN

    # Right Column: 2 Stacked Screenshots (Chat & Mentor Earnings)
    img_left = Inches(6.65)
    img_w = Inches(5.88)
    img_h1 = Inches(2.62)
    img_h2 = Inches(2.62)

    add_image_bordered(s, "08_chat.png", img_left, top_y, img_w, img_h1)
    add_image_bordered(s, "06b_mentor_earnings.png", img_left, top_y + img_h1 + Inches(0.16), img_w, img_h2)

    add_speaker_note(s,
        "Module 4 integrates Real-Time Messaging and the Payment Wallet. Our chat is powered by Socket.io for instantaneous delivery. "
        "To stop spam, non-paying users have a 5-message free limit, which converts to unlimited chat once a session is confirmed. "
        "Our wallet features a dual-balance system separating top-up funds from earned teaching revenue, with transparent 0% platform fees."
    )

# =================================================================================================
# SLIDE 8: Module 05: Role Dashboards, Analytics & Community Trust
# =================================================================================================
def build_slide_8():
    s = prs.slides.add_slide(blank_layout)
    set_slide_bg(s)
    add_header(s, "Module 05: Role Dashboards, Analytics & Community Trust", "CORE FUNCTIONAL MODULE")
    add_footer(s, 8)

    left_w = Inches(5.60)
    top_y = Inches(1.45)

    # Left Column: Features Card
    card_h = Inches(5.40)
    add_card(s, Inches(0.8), top_y, left_w, card_h, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb = s.shapes.add_textbox(Inches(1.05), top_y + Inches(0.20), left_w - Inches(0.5), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    ph = tf.paragraphs[0]
    ph.text = "Role-Tailored Dashboards & Feedback"
    ph.font.name = "Cambria"
    ph.font.size = Pt(14)
    ph.font.bold = True
    ph.font.color.rgb = DARK_GREEN
    ph.space_after = Pt(8)

    features = [
        ("Learner Mastery Workspace", "Visualizes learning momentum with key metrics: Sessions Attended, Hours of Live Learning, Mentors Connected, and Skills Explored."),
        ("Milestone Progress Bars", "Tracks active curriculum targets (e.g. Node.js 90%, System Design 72%) to keep students motivated toward their academic goals."),
        ("Mentor Teaching Studio", "Comprehensive overview of learners taught, sessions completed, average star rating, and granular skill depth across listed technologies."),
        ("Community Star Ratings & Reviews", "One-review-per-completed-session policy ensures 100% verified student testimonials and eliminates fake feedback."),
        ("Trust & Reputation Cycle", "Positive student reviews boost mentor rankings in the search directory, driving a self-sustaining cycle of quality education.")
    ]
    for hd, desc in features:
        p = tf.add_paragraph()
        p.space_after = Pt(6)
        r0 = p.add_run()
        r0.text = "• " + hd + ": "
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = TEXT_DARK
        r1 = p.add_run()
        r1.text = desc
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = TEXT_BODY

    p_tech = tf.add_paragraph()
    p_tech.space_before = Pt(8)
    p_tech.text = "APIs: GET /api/dashboard/learner, GET /api/dashboard/mentor, POST /api/reviews | Model: ReviewSchema"
    p_tech.font.name = "Calibri"
    p_tech.font.size = Pt(9)
    p_tech.font.bold = True
    p_tech.font.color.rgb = DARK_GREEN

    # Right Column: 2 Stacked Screenshots (Learner & Mentor Dashboards)
    img_left = Inches(6.65)
    img_w = Inches(5.88)
    img_h1 = Inches(2.62)
    img_h2 = Inches(2.62)

    add_image_bordered(s, "05_learner_dashboard.png", img_left, top_y, img_w, img_h1)
    add_image_bordered(s, "06_mentor_dashboard.png", img_left, top_y + img_h1 + Inches(0.16), img_w, img_h2)

    add_speaker_note(s,
        "Module 5 covers Role-Based Dashboards and Reviews. Learners have a dedicated workspace showing session hours and milestone bars. "
        "Mentors have a Teaching Studio summarizing learners reached and revenue earned. After each completed session, verified 1-5 star reviews "
        "can be submitted, building long-term academic reputation."
    )

# =================================================================================================
# SLIDE 9: Module 06: Admin Governance & System Verification
# =================================================================================================
def build_slide_9():
    s = prs.slides.add_slide(blank_layout)
    set_slide_bg(s)
    add_header(s, "Module 06: Admin Governance & System Verification", "GOVERNANCE & TESTING")
    add_footer(s, 9)

    left_w = Inches(5.60)
    top_y = Inches(1.45)

    # Left Column: Features Card
    card_h = Inches(5.40)
    add_card(s, Inches(0.8), top_y, left_w, card_h, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb = s.shapes.add_textbox(Inches(1.05), top_y + Inches(0.20), left_w - Inches(0.5), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    ph = tf.paragraphs[0]
    ph.text = "Platform Oversight & Quality Assurance"
    ph.font.name = "Cambria"
    ph.font.size = Pt(14)
    ph.font.bold = True
    ph.font.color.rgb = DARK_GREEN
    ph.space_after = Pt(8)

    features = [
        ("Institutional Overview KPIs", "Live executive counters tracking Total Users, Verified Mentors, Active Learners, Pending Verifications, and Total Sessions."),
        ("Member Moderation Controls", "Administrators can review profiles, approve unverified accounts, suspend bad actors, or delete fraudulent accounts."),
        ("Auditable Activity Logs", "All critical actions (role updates, booking cancellations, financial settlements) are stored in immutable MongoDB activity logs."),
        ("Comprehensive Testing Suite", "End-to-end user flows validated with automated Playwright browser tests; backend endpoints verified via Postman test collections."),
        ("Responsive & Dark Mode UI", "Seamless CSS theming tested across light/dark contrast standards and fully responsive down to 390px mobile viewports.")
    ]
    for hd, desc in features:
        p = tf.add_paragraph()
        p.space_after = Pt(6)
        r0 = p.add_run()
        r0.text = "• " + hd + ": "
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = TEXT_DARK
        r1 = p.add_run()
        r1.text = desc
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = TEXT_BODY

    p_tech = tf.add_paragraph()
    p_tech.space_before = Pt(8)
    p_tech.text = "APIs: GET /api/admin/stats, PUT /api/admin/users/:id/status | Security: AdminRBACGuard"
    p_tech.font.name = "Calibri"
    p_tech.font.size = Pt(9)
    p_tech.font.bold = True
    p_tech.font.color.rgb = DARK_GREEN

    # Right Column: 2 Stacked Screenshots (Admin Overview & Admin Members)
    img_left = Inches(6.65)
    img_w = Inches(5.88)
    img_h1 = Inches(2.62)
    img_h2 = Inches(2.62)

    add_image_bordered(s, "10_admin_dashboard.png", img_left, top_y, img_w, img_h1)
    add_image_bordered(s, "10b_admin_members.png", img_left, top_y + img_h1 + Inches(0.16), img_w, img_h2)

    add_speaker_note(s,
        "Module 6 covers Admin Governance and System Testing. Administrators have high-level oversight through live KPIs and can verify, "
        "suspend, or moderate community accounts. The entire platform has been thoroughly verified using Playwright automated browser tests "
        "and API validation suites to guarantee zero double-booking and resilient connection handling."
    )

# =================================================================================================
# SLIDE 10: Conclusion, Future Scope & Project Links (Exact UI as user's reference)
# =================================================================================================
def build_slide_10():
    s = prs.slides.add_slide(blank_layout)
    set_slide_bg(s, LIGHT_BG)

    # Top Bar: Logo + TEAM COSMICCODERS • GROUP 13
    logo_path = os.path.join(SCREENSHOTS_DIR, "logo.png")
    if os.path.exists(logo_path):
        s.shapes.add_picture(logo_path, Inches(0.8), Inches(0.7), Inches(0.65), Inches(0.65))

    tb_brand = s.shapes.add_textbox(Inches(1.55), Inches(0.65), Inches(4.5), Inches(0.7))
    tf_b = tb_brand.text_frame
    p_b1 = tf_b.paragraphs[0]
    p_b1.text = "SkillExchange"
    p_b1.font.name = "Cambria"
    p_b1.font.size = Pt(22)
    p_b1.font.bold = True
    p_b1.font.color.rgb = DARK_GREEN

    p_b2 = tf_b.add_paragraph()
    p_b2.text = "TEAM COSMICCODERS  •  GROUP 13"
    p_b2.font.name = "Calibri"
    p_b2.font.size = Pt(10)
    p_b2.font.bold = True
    p_b2.font.color.rgb = TEXT_MUTED

    # Left Column: Heading & Conclusion Points
    left_w = Inches(5.8)
    tb_left = s.shapes.add_textbox(Inches(0.8), Inches(1.55), left_w, Inches(5.0))
    tf_l = tb_left.text_frame
    tf_l.word_wrap = True

    p_th = tf_l.paragraphs[0]
    p_th.text = "Thank You!"
    p_th.font.name = "Cambria"
    p_th.font.size = Pt(40)
    p_th.font.bold = True
    p_th.font.color.rgb = DARK_GREEN
    p_th.space_after = Pt(2)

    p_qd = tf_l.add_paragraph()
    p_qd.text = "Questions & Technical Discussion"
    p_qd.font.name = "Calibri"
    p_qd.font.size = Pt(16)
    p_qd.font.color.rgb = TEXT_DARK
    p_qd.space_after = Pt(14)

    concl_pts = [
        ("Working Full-Stack Solution", "Engineered and deployed a functional MERN platform solving campus peer tutoring friction."),
        ("Zero Scheduling Friction", "Integrated Socket.io chat and deterministic 5-stage calendar scheduling with 4-hour auto-expiration."),
        ("In-Browser WebRTC Calling", "Future enhancement: Native peer-to-peer audio/video calling without external Google Meet links."),
        ("Automated Payouts & Alerts", "Future roadmap: Razorpay Route automated disbursements and browser push reminders 15 mins prior to sessions.")
    ]
    for ct, cd in concl_pts:
        p = tf_l.add_paragraph()
        p.space_after = Pt(6)
        r0 = p.add_run()
        r0.text = "• " + ct + ": "
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = DARK_GREEN
        r1 = p.add_run()
        r1.text = cd
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = TEXT_BODY

    # Right Column: Clean White Card for Project Links & Repositories
    card_w = Inches(5.6)
    card_h = Inches(4.8)
    card_x = Inches(6.9)
    card_y = Inches(1.55)
    add_card(s, card_x, card_y, card_w, card_h, bg_color=CARD_BG, border_color=BORDER_COLOR)

    tb_card = s.shapes.add_textbox(card_x + Inches(0.4), card_y + Inches(0.35), card_w - Inches(0.8), card_h - Inches(0.7))
    tf_c = tb_card.text_frame
    tf_c.word_wrap = True

    pc0 = tf_c.paragraphs[0]
    pc0.text = "PROJECT REPOSITORY & DEMO"
    pc0.font.name = "Calibri"
    pc0.font.size = Pt(11)
    pc0.font.bold = True
    pc0.font.color.rgb = DARK_GREEN
    pc0.space_after = Pt(12)

    pc_git_label = tf_c.add_paragraph()
    pc_git_label.text = "GitHub Source Code:"
    pc_git_label.font.name = "Calibri"
    pc_git_label.font.size = Pt(10.5)
    pc_git_label.font.bold = True
    pc_git_label.font.color.rgb = TEXT_DARK

    pc_git_val = tf_c.add_paragraph()
    pc_git_val.text = "github.com/CosmicCoder06/Skill-Exchange"
    pc_git_val.font.name = "Calibri"
    pc_git_val.font.size = Pt(11)
    pc_git_val.font.color.rgb = DARK_GREEN
    pc_git_val.space_after = Pt(14)

    pc_live_label = tf_c.add_paragraph()
    pc_live_label.text = "Live Cloud Deployment:"
    pc_live_label.font.name = "Calibri"
    pc_live_label.font.size = Pt(10.5)
    pc_live_label.font.bold = True
    pc_live_label.font.color.rgb = TEXT_DARK

    pc_live_val = tf_c.add_paragraph()
    pc_live_val.text = "skill-exchange-one-eta.vercel.app"
    pc_live_val.font.name = "Calibri"
    pc_live_val.font.size = Pt(11)
    pc_live_val.font.color.rgb = DARK_GREEN
    pc_live_val.space_after = Pt(16)

    pc_meta = tf_c.add_paragraph()
    pc_meta.text = "GLA University, Mathura  •  B.Tech CSE Capstone Project"
    pc_meta.font.name = "Calibri"
    pc_meta.font.size = Pt(10)
    pc_meta.font.color.rgb = TEXT_MUTED

    pc_meta2 = tf_c.add_paragraph()
    pc_meta2.text = "Internship Project at Codroid Hub Pvt. Ltd. (October 2026)"
    pc_meta2.font.name = "Calibri"
    pc_meta2.font.size = Pt(10)
    pc_meta2.font.color.rgb = TEXT_MUTED

    # Bottom Tagline & Slide Number
    tb_bot = s.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(8.0), Inches(0.3))
    tf_bot = tb_bot.text_frame
    p_bot = tf_bot.paragraphs[0]
    p_bot.text = "Learn together.  Share knowledge.  Grow together."
    p_bot.font.name = "Calibri"
    p_bot.font.size = Pt(11)
    p_bot.font.italic = True
    p_bot.font.color.rgb = TEXT_MUTED

    tb_num = s.shapes.add_textbox(Inches(11.5), Inches(6.95), Inches(1.0), Inches(0.3))
    tf_num = tb_num.text_frame
    p_num = tf_num.paragraphs[0]
    p_num.text = "10"
    p_num.font.name = "Calibri"
    p_num.font.size = Pt(14)
    p_num.font.bold = True
    p_num.font.color.rgb = DARK_GREEN
    p_num.alignment = PP_ALIGN.RIGHT

    add_speaker_note(s,
        "In conclusion, SkillExchange delivers an accessible, robust MERN solution for student mentorship. "
        "We have solved coordination ghosting and provided transparency in scheduling and peer honorariums. "
        "Our live demo is hosted on Vercel and the code is open on GitHub. We are now open for questions and feedback."
    )

def main():
    print("Building 10-slide presentation...")
    build_slide_1()
    build_slide_2()
    build_slide_3()
    build_slide_4()
    build_slide_5()
    build_slide_6()
    build_slide_7()
    build_slide_8()
    build_slide_9()
    build_slide_10()

    prs.save(OUTPUT_FILE)
    print(f"Presentation saved successfully to {OUTPUT_FILE} (Total slides: {len(prs.slides)})")

if __name__ == "__main__":
    main()
