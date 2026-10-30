import os
import re
from pptx import Presentation

INPUT_PATH = "Skill_Exchange_Capstone_Presentation.pptx"
OUTPUT_PATH = "Skill_Exchange_Capstone_Presentation_Simple.pptx"

def set_para_text(p, text):
    """Sets paragraph text while preserving font properties of the first run."""
    if not p.runs:
        p.text = text
        return
    
    r0 = p.runs[0]
    font_name = r0.font.name
    font_size = r0.font.size
    font_bold = r0.font.bold
    font_italic = r0.font.italic
    font_color_rgb = r0.font.color.rgb if (r0.font.color and hasattr(r0.font.color, 'rgb') and r0.font.color.rgb) else None
    
    # Empty subsequent runs
    for r in p.runs[1:]:
        r.text = ""
        
    r0.text = text
    if font_name:
        r0.font.name = font_name
    if font_size:
        r0.font.size = font_size
    if font_bold is not None:
        r0.font.bold = font_bold
    if font_italic is not None:
        r0.font.italic = font_italic
    if font_color_rgb:
        r0.font.color.rgb = font_color_rgb

def set_slide_note(slide, note_text):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = note_text

def simplify_presentation():
    prs = Presentation(INPUT_PATH)
    print(f"Loaded {INPUT_PATH} with {len(prs.slides)} slides.")
    
    # -------------------------------------------------------------
    # SLIDE 1: Title (Check and preserve empty fields)
    # -------------------------------------------------------------
    # Shape 2: Title & Details
    # Fields stay empty with underscores as requested.
    
    # -------------------------------------------------------------
    # SLIDE 2: Presentation Agenda
    # -------------------------------------------------------------
    s2 = prs.slides[1]
    # Shape 0: Tag
    set_para_text(s2.shapes[0].text_frame.paragraphs[0], "OVERVIEW")
    # Shape 1: Title
    set_para_text(s2.shapes[1].text_frame.paragraphs[0], "Presentation Agenda")
    # Card 01: Shape 6
    set_para_text(s2.shapes[6].text_frame.paragraphs[0], "Problem & Objectives")
    set_para_text(s2.shapes[6].text_frame.paragraphs[1], "Problems students face when finding study help, and what our project aims to do.")
    # Card 02: Shape 9
    set_para_text(s2.shapes[9].text_frame.paragraphs[0], "Design & Database")
    set_para_text(s2.shapes[9].text_frame.paragraphs[1], "Full-stack MERN stack, real-time chat with WebSockets, and MongoDB data models.")
    # Card 03: Shape 12
    set_para_text(s2.shapes[12].text_frame.paragraphs[0], "Main feature: Session Booking in 5 steps")
    set_para_text(s2.shapes[12].text_frame.paragraphs[1], "5-step booking system with automatic 4-hour reply timer.")
    # Card 04: Shape 15
    set_para_text(s2.shapes[15].text_frame.paragraphs[0], "Live System Walkthrough")
    set_para_text(s2.shapes[15].text_frame.paragraphs[1], "Live demo of Home, Login, Profiles, Search, Dashboards, Chat, and Admin screens.")
    # Card 05: Shape 18
    set_para_text(s2.shapes[18].text_frame.paragraphs[0], "Engineering & Future Scope")
    set_para_text(s2.shapes[18].text_frame.paragraphs[1], "A CSS dark theme bug we fixed, current project limits, and future plans.")
    
    set_slide_note(s2, 
        "Good morning teachers. Today we are presenting our capstone project called Skill Exchange. "
        "On this slide, you can see what we will cover today. First, we will discuss the problems students face in peer learning. "
        "Then, we will explain our tech stack, system design, booking flow, live screens, and future scope."
    )

    # -------------------------------------------------------------
    # SLIDE 3: Problem Statement: Challenges in Peer Learning
    # -------------------------------------------------------------
    s3 = prs.slides[2]
    set_para_text(s3.shapes[0].text_frame.paragraphs[0], "CONTEXT & BACKGROUND")
    set_para_text(s3.shapes[1].text_frame.paragraphs[0], "Problem Statement: Challenges in Peer Learning")
    # Stat 1: Shape 5
    set_para_text(s3.shapes[5].text_frame.paragraphs[0], "TOO COSTLY")
    set_para_text(s3.shapes[5].text_frame.paragraphs[1], "Commercial 1-on-1 tutoring charges steep hourly rates that normal college students cannot afford.")
    # Stat 2: Shape 7
    set_para_text(s3.shapes[7].text_frame.paragraphs[0], "SCATTERED")
    set_para_text(s3.shapes[7].text_frame.paragraphs[1], "Student study talks are spread across WhatsApp, Discord, and Google Calendar.")
    # Stat 3: Shape 9
    set_para_text(s3.shapes[9].text_frame.paragraphs[0], "NO COMMITMENT")
    set_para_text(s3.shapes[9].text_frame.paragraphs[1], "Casual peer study often leads to no-shows and people cancelling at the last minute.")
    # Detail 1: Shape 11
    set_para_text(s3.shapes[11].text_frame.paragraphs[0], "1. Expensive Private Tutoring")
    set_para_text(s3.shapes[11].text_frame.paragraphs[1], "Big learning apps focus on high-paying users, charging fees that college students cannot afford.")
    # Detail 2: Shape 13
    set_para_text(s3.shapes[13].text_frame.paragraphs[0], "2. Too many different apps")
    set_para_text(s3.shapes[13].text_frame.paragraphs[1], "Setting up study calls over chat groups leads to missed links, wrong timings, and lost notes.")
    # Detail 3: Shape 15
    set_para_text(s3.shapes[15].text_frame.paragraphs[0], "3. High Ghosting & No-Show Rates")
    set_para_text(s3.shapes[15].text_frame.paragraphs[1], "Free peer help has no fixed schedule, so people often cancel at the last minute.")
    # Detail 4: Shape 17
    set_para_text(s3.shapes[17].text_frame.paragraphs[0], "4. Money talk is unclear")
    set_para_text(s3.shapes[17].text_frame.paragraphs[1], "Students feel shy to ask peers for money, and having no clear fee record causes confusion.")

    set_slide_note(s3,
        "When students need help with difficult subjects, they face four main problems. "
        "Private tutoring is too expensive for most college students. "
        "Coordinating sessions across WhatsApp and Discord is messy and unorganized. "
        "Many times students cancel last minute because there is no formal commitment. "
        "Also, talking about session fees with peers feels awkward, so our app solves this."
    )

    # -------------------------------------------------------------
    # SLIDE 4: Project Objectives & Core Modules
    # -------------------------------------------------------------
    s4 = prs.slides[3]
    set_para_text(s4.shapes[0].text_frame.paragraphs[0], "SCOPE & ROLES")
    set_para_text(s4.shapes[1].text_frame.paragraphs[0], "Project Objectives & Core Modules")
    # Banner: Shape 5
    set_para_text(s4.shapes[5].text_frame.paragraphs[0], "Primary Goal: A trusted campus platform for affordable 1-on-1 peer learning.")
    set_para_text(s4.shapes[5].text_frame.paragraphs[1], "Help skilled students earn by teaching, while letting learners book organized, 1-on-1 study sessions.")
    # Learner Card: Shape 7
    set_para_text(s4.shapes[7].text_frame.paragraphs[0], "Learner Module")
    set_para_text(s4.shapes[7].text_frame.paragraphs[1], "For Students Seeking Guidance")
    set_para_text(s4.shapes[7].text_frame.paragraphs[2], "• Discover verified mentors by skill tags & hourly rate")
    set_para_text(s4.shapes[7].text_frame.paragraphs[3], "• 3-step session scheduling with date and time slots")
    set_para_text(s4.shapes[7].text_frame.paragraphs[4], "• Real-time 1-on-1 chat (5 free messages before booking)")
    set_para_text(s4.shapes[7].text_frame.paragraphs[5], "• Track total learning hours, attended sessions & progress")
    # Mentor Card: Shape 9
    set_para_text(s4.shapes[9].text_frame.paragraphs[0], "Mentor Module")
    set_para_text(s4.shapes[9].text_frame.paragraphs[1], "For Student Tutors")
    set_para_text(s4.shapes[9].text_frame.paragraphs[2], "• Set custom hourly fees and select free time slots")
    set_para_text(s4.shapes[9].text_frame.paragraphs[3], "• Showcase skills with levels (Beginner to Expert)")
    set_para_text(s4.shapes[9].text_frame.paragraphs[4], "• Accept, decline, or suggest another session time")
    set_para_text(s4.shapes[9].text_frame.paragraphs[5], "• Teaching dashboard to see students and track earnings")
    # Admin Card: Shape 11
    set_para_text(s4.shapes[11].text_frame.paragraphs[0], "Admin Control")
    set_para_text(s4.shapes[11].text_frame.paragraphs[1], "For Platform Safety & Verification")
    set_para_text(s4.shapes[11].text_frame.paragraphs[2], "• Central dashboard: see user counts, sessions, and pending checks")
    set_para_text(s4.shapes[11].text_frame.paragraphs[3], "• User management: verify student accounts or block bad users")
    set_para_text(s4.shapes[11].text_frame.paragraphs[4], "• Check booking status and see user reports")
    set_para_text(s4.shapes[11].text_frame.paragraphs[5], "• Download payment records as CSV files")

    set_slide_note(s4,
        "Our main goal is to make peer tutoring simple and organized inside a college. "
        "We built three user roles: Learner, Mentor, and Admin. "
        "Learners can search for mentors, book a slot, and chat with them. "
        "Mentors can set their own hourly rates, choose when they are free, and track earnings. "
        "The admin oversees the whole platform, verifies users, and can download payment records."
    )

    # -------------------------------------------------------------
    # SLIDE 5: Technology Stack & Database Design
    # -------------------------------------------------------------
    s5 = prs.slides[4]
    set_para_text(s5.shapes[0].text_frame.paragraphs[0], "ENGINEERING FOUNDATION")
    set_para_text(s5.shapes[1].text_frame.paragraphs[0], "Technology Stack & Database Design")
    # Shape 5: Frontend
    set_para_text(s5.shapes[5].text_frame.paragraphs[0], "Frontend: React 19 + Vite")
    set_para_text(s5.shapes[5].text_frame.paragraphs[1], "Component-based UI, fast page routing, Lucide icons, and clean CSS styling.")
    # Shape 7: Backend
    set_para_text(s5.shapes[7].text_frame.paragraphs[0], "Backend: Node.js + Express")
    set_para_text(s5.shapes[7].text_frame.paragraphs[1], "REST API controllers, clean routes, error handling, and input checks.")
    # Shape 9: Database
    set_para_text(s5.shapes[9].text_frame.paragraphs[0], "Database: MongoDB Atlas")
    set_para_text(s5.shapes[9].text_frame.paragraphs[1], "Cloud NoSQL database with Mongoose schemas for data relationships.")
    # Shape 11: Real-time
    set_para_text(s5.shapes[11].text_frame.paragraphs[0], "Real-Time: Socket.io")
    set_para_text(s5.shapes[11].text_frame.paragraphs[1], "Live two-way chat using Socket.io for instant messaging.")
    # Shape 13: Auth
    set_para_text(s5.shapes[13].text_frame.paragraphs[0], "Auth: JWT & Bcrypt")
    set_para_text(s5.shapes[13].text_frame.paragraphs[1], "JWT: token that proves user login; Bcrypt: safely scrambles passwords.")
    # Shape 15: Payment & Wallet
    set_para_text(s5.shapes[15].text_frame.paragraphs[0], "Payment & Wallet")
    set_para_text(s5.shapes[15].text_frame.paragraphs[1], "0% platform fee, 0% tax, test wallet balance, and manual payouts by admin.")
    # Shape 17: User Model
    set_para_text(s5.shapes[17].text_frame.paragraphs[0], "User Model")
    set_para_text(s5.shapes[17].text_frame.paragraphs[1], "name, email, password, role ('learner'|'mentor'|'admin'), bio, skillsToTeach, skillsToLearn, hourlyRate, availability, payoutUpiId, profileCompleted, isActive")
    # Shape 19: Booking Model
    set_para_text(s5.shapes[19].text_frame.paragraphs[0], "Booking Model")
    set_para_text(s5.shapes[19].text_frame.paragraphs[1], "learner (Ref), mentor (Ref), date, time, duration (30-120m), status: pending / confirmed / completed / cancelled, actionExpiresAt (4h), meetingUrl")
    # Shape 21: Chat Models
    set_para_text(s5.shapes[21].text_frame.paragraphs[0], "Chat Models")
    set_para_text(s5.shapes[21].text_frame.paragraphs[1], "Conversation: participants, bookingId, lastMessage, lastActivityAt || Message: conversationId, sender, text, isEdited, readBy")
    # Shape 23: Review & Wallet
    set_para_text(s5.shapes[23].text_frame.paragraphs[0], "Review & Wallet Models")
    set_para_text(s5.shapes[23].text_frame.paragraphs[1], "Review: mentor, learner, rating (1-5), comment || Wallet: balance (test top-up), user, transactions list")

    set_slide_note(s5,
        "Here is the technology behind our application. We used the MERN stack with React 19 on the frontend and Node.js with Express on the backend. "
        "For the database, we use MongoDB Atlas with Mongoose. Chat works instantly using Socket.io. "
        "For security, passwords are encrypted using Bcrypt and users get a secure JWT login token. "
        "The bottom boxes show our main database models."
    )

    # -------------------------------------------------------------
    # SLIDE 6: System Architecture: End-to-End Data Flow
    # -------------------------------------------------------------
    s6 = prs.slides[5]
    set_para_text(s6.shapes[0].text_frame.paragraphs[0], "HIGH-LEVEL DESIGN")
    set_para_text(s6.shapes[1].text_frame.paragraphs[0], "System Architecture: End-to-End Data Flow")
    # Shape 4: React Client
    set_para_text(s6.shapes[4].text_frame.paragraphs[0], "React Client (Frontend)")
    set_para_text(s6.shapes[4].text_frame.paragraphs[1], "• Fast page rendering with Vite")
    set_para_text(s6.shapes[4].text_frame.paragraphs[2], "• Page switching with React Router")
    set_para_text(s6.shapes[4].text_frame.paragraphs[3], "• Separate Learner & Mentor screens")
    set_para_text(s6.shapes[4].text_frame.paragraphs[4], "• Real-time chat listener with Socket.io")
    set_para_text(s6.shapes[4].text_frame.paragraphs[5], "• Stores login token in local storage")
    # Shape 6: REST API Calls
    set_para_text(s6.shapes[6].text_frame.paragraphs[0], "REST API Calls")
    # Shape 8: WebSocket (Chat)
    set_para_text(s6.shapes[8].text_frame.paragraphs[0], "WebSocket (Live Chat)")
    # Shape 9: Node.js / Express API
    set_para_text(s6.shapes[9].text_frame.paragraphs[0], "Node.js / Express API")
    set_para_text(s6.shapes[9].text_frame.paragraphs[1], "• JWT login token check & user roles")
    set_para_text(s6.shapes[9].text_frame.paragraphs[2], "• Session booking status controller")
    set_para_text(s6.shapes[9].text_frame.paragraphs[3], "• Socket.io server for chat messages")
    set_para_text(s6.shapes[9].text_frame.paragraphs[4], "• Creates Google Meet call link")
    set_para_text(s6.shapes[9].text_frame.paragraphs[5], "• Admin control routes")
    # Shape 11: Mongoose Driver
    set_para_text(s6.shapes[11].text_frame.paragraphs[0], "Mongoose Driver")
    # Shape 12: MongoDB Atlas
    set_para_text(s6.shapes[12].text_frame.paragraphs[0], "MongoDB Atlas")
    set_para_text(s6.shapes[12].text_frame.paragraphs[1], "• Users")
    set_para_text(s6.shapes[12].text_frame.paragraphs[2], "• Bookings")
    set_para_text(s6.shapes[12].text_frame.paragraphs[3], "• Messages")
    set_para_text(s6.shapes[12].text_frame.paragraphs[4], "• Reviews")
    set_para_text(s6.shapes[12].text_frame.paragraphs[5], "• Activity Logs")

    set_slide_note(s6,
        "This diagram shows how data moves across our application. "
        "On the left, the React frontend talks to the backend using standard REST API calls for pages, and WebSockets for live chat. "
        "In the center, our Express server checks user login tokens and handles booking logic. "
        "On the right, the server reads and saves all data into MongoDB Atlas collections."
    )

    # -------------------------------------------------------------
    # SLIDE 7: Signature Booking Flow -> Main feature: Session Booking in 5 steps
    # -------------------------------------------------------------
    s7 = prs.slides[6]
    set_para_text(s7.shapes[0].text_frame.paragraphs[0], "MAIN FEATURE")
    set_para_text(s7.shapes[1].text_frame.paragraphs[0], "Main feature: Session Booking in 5 steps")
    # Step 1: Shape 5
    set_para_text(s7.shapes[5].text_frame.paragraphs[0], "STEP 1")
    set_para_text(s7.shapes[5].text_frame.paragraphs[1], "Choose")
    set_para_text(s7.shapes[5].text_frame.paragraphs[2], "Learner picks session length (30-120 min), selects a date, and chooses a time slot.")
    # Step 2: Shape 8
    set_para_text(s7.shapes[8].text_frame.paragraphs[0], "STEP 2")
    set_para_text(s7.shapes[8].text_frame.paragraphs[1], "Check")
    set_para_text(s7.shapes[8].text_frame.paragraphs[2], "System checks mentor schedule, blocks self-booking, and sets 4-hour reply timer.")
    # Step 3: Shape 11
    set_para_text(s7.shapes[11].text_frame.paragraphs[0], "STEP 3")
    set_para_text(s7.shapes[11].text_frame.paragraphs[1], "Reply")
    set_para_text(s7.shapes[11].text_frame.paragraphs[2], "Mentor sees the request and can Accept, Suggest Another Time, or Decline.")
    # Step 4: Shape 14
    set_para_text(s7.shapes[14].text_frame.paragraphs[0], "STEP 4")
    set_para_text(s7.shapes[14].text_frame.paragraphs[1], "Confirm")
    set_para_text(s7.shapes[14].text_frame.paragraphs[2], "Session is confirmed, Google Meet link is added, and unlimited chat is unlocked.")
    # Step 5: Shape 17
    set_para_text(s7.shapes[17].text_frame.paragraphs[0], "STEP 5")
    set_para_text(s7.shapes[17].text_frame.paragraphs[1], "Complete")
    set_para_text(s7.shapes[17].text_frame.paragraphs[2], "Session takes place, earnings update in wallet, and both give ratings.")
    # Banner: Shape 19
    set_para_text(s7.shapes[19].text_frame.paragraphs[0], "4-hour reply time and other time suggestion")
    set_para_text(s7.shapes[19].text_frame.paragraphs[1], "• If a mentor does not reply within 4 hours, the request cancels automatically so the learner can find another mentor.")
    set_para_text(s7.shapes[19].text_frame.paragraphs[2], "• Mentors can suggest up to 3 alternate times instead of directly rejecting the learner.")

    set_slide_note(s7,
        "This slide explains our 5-step booking system, which is the core feature of our app. "
        "First, the learner selects the duration, date, and time. "
        "Second, the system validates the slot and sets a 4-hour timer. "
        "Third, the mentor can accept, decline, or suggest other times. Once accepted, a Google Meet link is provided and chat unlocks. "
        "Finally, after the session, both users can leave reviews."
    )

    # -------------------------------------------------------------
    # SLIDE 8: Live Walkthrough: Landing & Authentication
    # -------------------------------------------------------------
    s8 = prs.slides[7]
    set_para_text(s8.shapes[0].text_frame.paragraphs[0], "USER ONBOARDING")
    set_para_text(s8.shapes[1].text_frame.paragraphs[0], "Live Walkthrough: Landing & Authentication")
    set_para_text(s8.shapes[4].text_frame.paragraphs[0], "• Two-Path Landing Page: Gives clear buttons for 'Share what you know' (mentor) and 'Learn something new' (learner).")
    set_para_text(s8.shapes[4].text_frame.paragraphs[1], "• Role-Based Signup: Users pick whether they want to 'Learn Skills' or 'Teach & Mentor' during registration.")
    set_para_text(s8.shapes[4].text_frame.paragraphs[2], "• Safe Login: Uses Bcrypt password encryption, stores JWT token in browser, and opens the correct dashboard.")

    set_slide_note(s8,
        "These screenshots show our landing and authentication pages. "
        "The home page welcomes visitors and gives them a clear choice to learn or teach. "
        "On the signup page, students select their role as a Learner or Mentor. "
        "When they log in, their password is encrypted safely and they are redirected to their personalized dashboard."
    )

    # -------------------------------------------------------------
    # SLIDE 9: Live Walkthrough: Discovery & Profile Management
    # -------------------------------------------------------------
    s9 = prs.slides[8]
    set_para_text(s9.shapes[0].text_frame.paragraphs[0], "PEER MATCHING")
    set_para_text(s9.shapes[1].text_frame.paragraphs[0], "Live Walkthrough: Discovery & Profile Management")
    set_para_text(s9.shapes[4].text_frame.paragraphs[0], "• Searchable Directory: Search mentors and learners instantly by name, role, or technical skill tags.")
    set_para_text(s9.shapes[4].text_frame.paragraphs[1], "• Detailed Profile: Shows bio, skills to teach, skills to learn, and weekly free time slots.")
    set_para_text(s9.shapes[4].text_frame.paragraphs[2], "• Public Mentor Card: Shows hourly rate (e.g. ₹400/hr), top skills, and a direct 'Book Session' button.")

    set_slide_note(s9,
        "Here we show how learners find mentors. On the Discover page, students can search by typing names or filtering by skills like React or Python. "
        "The profile page shows the mentor's bio, ratings, hourly fee, and available hours. "
        "From here, a learner can click 'Book Session' to schedule a call directly."
    )

    # -------------------------------------------------------------
    # SLIDE 10: Live Walkthrough: Role-Based Dashboards
    # -------------------------------------------------------------
    s10 = prs.slides[9]
    set_para_text(s10.shapes[0].text_frame.paragraphs[0], "PROGRESS & EARNINGS")
    set_para_text(s10.shapes[1].text_frame.paragraphs[0], "Live Walkthrough: Role-Based Dashboards")
    set_para_text(s10.shapes[4].text_frame.paragraphs[0], "• Learner Dashboard: Shows total learning hours completed, upcoming sessions, and connected mentors.")
    set_para_text(s10.shapes[4].text_frame.paragraphs[1], "• Mentor Teaching Studio: Shows incoming session requests, student count, and skill level meters.")
    set_para_text(s10.shapes[4].text_frame.paragraphs[2], "• Earnings Record: Displays total money earned, pending requests, and UPI ID for payouts.")

    set_slide_note(s10,
        "Both learners and mentors have customized dashboards to track their progress. "
        "The learner dashboard displays hours learned, completed sessions, and next scheduled calls. "
        "The mentor dashboard shows session requests they need to accept or decline, student counts, and their total earnings along with their saved UPI ID."
    )

    # -------------------------------------------------------------
    # SLIDE 11: Live Walkthrough: Session Booking & Live Chat
    # -------------------------------------------------------------
    s11 = prs.slides[10]
    set_para_text(s11.shapes[0].text_frame.paragraphs[0], "CORE COLLABORATION")
    set_para_text(s11.shapes[1].text_frame.paragraphs[0], "Live Walkthrough: Session Booking & Live Chat")
    set_para_text(s11.shapes[4].text_frame.paragraphs[0], "• 3-Step Booking Window: Pick duration (30 to 120 mins), select an upcoming date, and choose an available time slot.")
    set_para_text(s11.shapes[4].text_frame.paragraphs[1], "• Clear Pricing: 0% platform fee, 0% tax, clear calculation based on minutes, with test wallet top-up.")
    set_para_text(s11.shapes[4].text_frame.paragraphs[2], "• Live two-way chat using Socket.io: Instant chat with a 5 free messages limit before booking to prevent spam.")

    set_slide_note(s11,
        "This slide shows the booking modal and the chat screen. "
        "Booking a session takes just three simple steps: pick duration, pick a date, and select a time. "
        "The pricing is completely transparent with zero platform fees. "
        "Our chat uses Socket.io for instant messaging, with a limit of 5 free messages before booking to stop spam."
    )

    # -------------------------------------------------------------
    # SLIDE 12: Live Walkthrough: Admin Control & Quality Checks
    # -------------------------------------------------------------
    s12 = prs.slides[11]
    set_para_text(s12.shapes[0].text_frame.paragraphs[0], "ADMIN CONTROL")
    set_para_text(s12.shapes[1].text_frame.paragraphs[0], "Live Walkthrough: Admin Control & Quality Checks")
    set_para_text(s12.shapes[4].text_frame.paragraphs[0], "• Platform Summary Numbers: Tracks total members (14), mentors (8), learners (5), and sessions (29).")
    set_para_text(s12.shapes[4].text_frame.paragraphs[1], "• Member Management: Admin can view student details, change user roles, or block bad users.")
    set_para_text(s12.shapes[4].text_frame.paragraphs[2], "• Reviews & Records: Shows 1 to 5 star ratings given after sessions and allows exporting payment records as CSV.")

    set_slide_note(s12,
        "Here is the admin panel, which helps manage the entire website. "
        "The admin can see live stats like total registered users, active mentors, and total sessions. "
        "The admin can also verify accounts, deactivate suspicious users, review ratings left by students, and download payment records into CSV spreadsheets."
    )

    # -------------------------------------------------------------
    # SLIDE 13: A Problem We Faced: Dark Theme Bug
    # -------------------------------------------------------------
    s13 = prs.slides[12]
    set_para_text(s13.shapes[0].text_frame.paragraphs[0], "TECHNICAL CHALLENGE")
    set_para_text(s13.shapes[1].text_frame.paragraphs[0], "A Problem We Faced: Dark Theme Bug")
    # Left Box: Shape 5
    set_para_text(s13.shapes[5].text_frame.paragraphs[0], "The Problem")
    set_para_text(s13.shapes[5].text_frame.paragraphs[1], "When we first added dark mode, the public pages worked fine, but the Admin sidebar and tables broke with washed-out white boxes and unreadable text.")
    set_para_text(s13.shapes[5].text_frame.paragraphs[2], "Why it happened")
    set_para_text(s13.shapes[5].text_frame.paragraphs[3], "High-priority CSS rules in the main stylesheet conflicted with admin styles, overriding dark backgrounds with inverted white colors.")
    set_para_text(s13.shapes[5].text_frame.paragraphs[4], "How we fixed it")
    set_para_text(s13.shapes[5].text_frame.paragraphs[5], "We properly scoped admin styles under html[data-theme='dark'] .admin-page and .admin-sidebar, making sure text and backgrounds have clear contrast.")
    # Right Box: Shape 7 (Dark code box - keep code intact, just simplify comment)
    set_para_text(s13.shapes[7].text_frame.paragraphs[1], "/* Keeps the dark green sidebar */")
    set_para_text(s13.shapes[7].text_frame.paragraphs[8], "/* Scoped admin page background & text */")

    set_slide_note(s13,
        "During development, we faced an interesting technical bug with dark mode. "
        "The student pages looked good, but the admin dashboard had conflicting CSS rules that made tables unreadable. "
        "We solved this by properly scoping admin CSS rules with specific theme selectors, ensuring consistent dark colors and clear contrast across all screens."
    )

    # -------------------------------------------------------------
    # SLIDE 14: Project Conclusion & Future Scope
    # -------------------------------------------------------------
    s14 = prs.slides[13]
    set_para_text(s14.shapes[0].text_frame.paragraphs[0], "SUMMARY & NEXT STEPS")
    set_para_text(s14.shapes[1].text_frame.paragraphs[0], "Project Conclusion & Future Scope")
    # Left Box: Shape 5
    set_para_text(s14.shapes[5].text_frame.paragraphs[0], "Key Project Achievements")
    set_para_text(s14.shapes[5].text_frame.paragraphs[1], "• Full MERN stack web app: Solved student peer learning issues using React 19, Express REST APIs, and MongoDB Atlas.")
    set_para_text(s14.shapes[5].text_frame.paragraphs[2], "• No platform fee: Created a fair learning platform with 0% commission and clear session tracking.")
    set_para_text(s14.shapes[5].text_frame.paragraphs[3], "• Reliable Session Booking: Automatic 4-hour cancellation and alternate time suggestions prevent ghosting.")
    set_para_text(s14.shapes[5].text_frame.paragraphs[4], "• Real-Time Chat & Meetings: Built-in instant Socket.io chat and automatic Google Meet video links for every session.")
    set_para_text(s14.shapes[5].text_frame.paragraphs[5], "• Made during our internship: Successfully built by Team CosmicCoders (Group 13) during our internship at Codroid Hub Pvt. Ltd.")
    # Right Box: Shape 7
    set_para_text(s14.shapes[7].text_frame.paragraphs[0], "Future Scope")
    set_para_text(s14.shapes[7].text_frame.paragraphs[1], "• In-Browser Video Calling:")
    set_para_text(s14.shapes[7].text_frame.paragraphs[2], "Add built-in video and audio calling using WebRTC directly inside the browser without needing Google Meet.")
    set_para_text(s14.shapes[7].text_frame.paragraphs[3], "• Online Payment Gateways:")
    set_para_text(s14.shapes[7].text_frame.paragraphs[4], "Integrate live payment gateways like Razorpay for real UPI payments and automatic refunds.")
    set_para_text(s14.shapes[7].text_frame.paragraphs[5], "• Push Notifications & Reminders:")
    set_para_text(s14.shapes[7].text_frame.paragraphs[6], "Add browser push notifications and calendar alerts 15 minutes before sessions start.")
    set_para_text(s14.shapes[7].text_frame.paragraphs[7], "• Mobile App:")
    set_para_text(s14.shapes[7].text_frame.paragraphs[8], "Build mobile apps for Android and iOS using React Native so students can chat and book on their phones.")

    set_slide_note(s14,
        "In conclusion, our team successfully built a complete MERN stack peer learning platform during our internship at Codroid Hub. "
        "It provides reliable session booking, zero platform fees, and real-time chat. "
        "For future scope, we plan to add in-browser WebRTC video calls, live Razorpay payments, calendar reminders, and mobile apps for Android and iOS."
    )

    # -------------------------------------------------------------
    # SLIDE 15: Thank You (Remove heart emoji)
    # -------------------------------------------------------------
    s15 = prs.slides[14]
    # Shape 2: Title & links
    set_para_text(s15.shapes[2].text_frame.paragraphs[4], "Built by Team CosmicCoders (Group 13)")

    prs.save(OUTPUT_PATH)
    print(f"Saved successfully to {OUTPUT_PATH}")

if __name__ == "__main__":
    simplify_presentation()
