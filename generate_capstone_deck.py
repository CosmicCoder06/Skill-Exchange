import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    DARK_GREEN = RGBColor(0x17, 0x6B, 0x4E)
    DARK_GREEN_DEEP = RGBColor(0x0F, 0x47, 0x33)
    LIGHT_MINT = RGBColor(0xF2, 0xF7, 0xF4)
    CARD_BG = RGBColor(0xF8, 0xFB, 0xF9)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    BLACK = RGBColor(0x00, 0x00, 0x00)
    DARK_TEXT = RGBColor(0x22, 0x22, 0x22)
    MUTED_TEXT = RGBColor(0x66, 0x66, 0x66)
    BORDER_COLOR = RGBColor(0xC5, 0xD0, 0xC9)
    CODE_BG = RGBColor(0x1E, 0x24, 0x22)
    CODE_GREEN = RGBColor(0x7E, 0xD4, 0xAD)
    TOTAL_SLIDES = 15

    base_dir = os.path.abspath(os.path.dirname(__file__))
    img_dir = os.path.join(base_dir, 'screenshots')

    def add_header(slide, title_text, category_tag=None):
        top_offset = Inches(0.4)
        if category_tag:
            tb_tag = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.7), Inches(0.3))
            tf_t = tb_tag.text_frame
            tf_t.word_wrap = True
            tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
            p_t = tf_t.paragraphs[0]
            p_t.text = category_tag.upper()
            p_t.font.name = "Calibri"
            p_t.font.size = Pt(10)
            p_t.font.bold = True
            p_t.font.color.rgb = DARK_GREEN
            top_offset = Inches(0.65)

        tx_box = slide.shapes.add_textbox(Inches(0.8), top_offset, Inches(11.7), Inches(0.6))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Cambria"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = DARK_GREEN

    def add_footer(slide, slide_num):
        # Left footer: Project name
        tb_l = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(6.0), Inches(0.3))
        tf_l = tb_l.text_frame
        tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0
        p_l = tf_l.paragraphs[0]
        p_l.text = "Skill Exchange · Capstone Project"
        p_l.font.name = "Calibri"
        p_l.font.size = Pt(10)
        p_l.font.color.rgb = MUTED_TEXT

        # Right footer: Slide number
        tb_r = slide.shapes.add_textbox(Inches(7.5), Inches(7.05), Inches(5.033), Inches(0.3))
        tf_r = tb_r.text_frame
        tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
        p_r = tf_r.paragraphs[0]
        p_r.text = f"Slide {slide_num} of {TOTAL_SLIDES}"
        p_r.font.name = "Calibri"
        p_r.font.size = Pt(10)
        p_r.font.color.rgb = MUTED_TEXT
        p_r.alignment = PP_ALIGN.RIGHT

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    def add_image_bordered(slide, filename, left, top, width=None, height=None):
        path = os.path.join(img_dir, filename)
        if not os.path.exists(path):
            print(f"Warning: {filename} missing!")
            return None
        pic = slide.shapes.add_picture(path, left, top, width=width, height=height)
        pic.line.color.rgb = BORDER_COLOR
        pic.line.width = Pt(1)
        return pic

    # ==========================================
    # SLIDE 1: Title Slide (Full Bleed Dark Green)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_GREEN
    bg1.line.fill.background()

    # Logo top-right
    logo_path = os.path.join(img_dir, "logo_white.png")
    if os.path.exists(logo_path):
        s1.shapes.add_picture(logo_path, Inches(11.2), Inches(0.8), width=Inches(1.2))

    tb1 = s1.shapes.add_textbox(Inches(0.9), Inches(0.9), Inches(10.0), Inches(5.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "CAPSTONE PROJECT · MERN STACK"
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xA5, 0xE0, 0xC8)
    p.space_after = Pt(8)

    p = tf1.add_paragraph()
    p.text = "Skill Exchange"
    p.font.name = "Cambria"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(4)

    p = tf1.add_paragraph()
    p.text = "A platform connecting mentors and learners"
    p.font.name = "Calibri"
    p.font.size = Pt(20)
    p.font.color.rgb = RGBColor(0xDF, 0xF0, 0xE8)
    p.space_after = Pt(20)

    p = tf1.add_paragraph()
    p.text = "Team CosmicCoders · Group 13 | B.Tech CSE, GLA University, Mathura"
    p.font.name = "Calibri"
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE
    p.space_after = Pt(4)

    p = tf1.add_paragraph()
    p.text = "Internship Project at Codroid Hub Pvt. Ltd. | October 2026"
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(0xCD, 0xE5, 0xDC)
    p.space_after = Pt(26)

    placeholders = [
        "Team Leader:   __________________________________________________",
        "Team Members: __________________________________________________",
        "Project Guide: __________________________________________________"
    ]
    for pl in placeholders:
        p = tf1.add_paragraph()
        p.text = pl
        p.font.name = "Calibri"
        p.font.size = Pt(14)
        p.font.color.rgb = WHITE
        p.space_after = Pt(10)

    # ==========================================
    # SLIDE 2: Agenda (5 Numbered Cards)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Presentation Agenda", "Overview")
    add_footer(s2, 2)

    agenda_items = [
        ("01", "Problem & Objectives", "Campus peer learning challenges, informal study friction, and core project goals."),
        ("02", "Architecture & Schemas", "Full-stack MERN ecosystem, bi-directional WebSockets, and Mongoose database design."),
        ("03", "Signature Booking Flow", "5-stage deterministic session lifecycle with automatic 4-hour expiration policy."),
        ("04", "Live System Walkthrough", "UI tour across Home, Auth, Profiles, Directory, Dashboards, Chat, and Admin Console."),
        ("05", "Engineering & Future Scope", "Dark theme CSS specificity resolution, technical constraints, and future roadmap.")
    ]

    card_y = Inches(1.4)
    card_h = Inches(0.95)
    for num, title, desc in agenda_items:
        add_card(s2, Inches(0.8), card_y, Inches(11.733), card_h, bg_color=CARD_BG)

        # Number pill on left
        pill = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), card_y + Inches(0.18), Inches(0.7), Inches(0.6))
        pill.fill.solid()
        pill.fill.fore_color.rgb = DARK_GREEN
        pill.line.fill.background()
        p = pill.text_frame.paragraphs[0]
        p.text = num
        p.font.name = "Cambria"
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

        # Title & Desc text
        tb = s2.shapes.add_textbox(Inches(1.9), card_y + Inches(0.12), Inches(10.3), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Cambria"
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = DARK_GREEN
        p1.space_after = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(13)
        p2.font.color.rgb = DARK_TEXT

        card_y += Inches(1.08)

    # ==========================================
    # SLIDE 3: Problem Statement (3 Stat Callouts + 4 Cards)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Problem Statement: Challenges in Peer Learning", "Context & Background")
    add_footer(s3, 3)

    # Top 3 Big Stat Callouts
    stats = [
        ("HIGH COST", "Commercial 1-on-1 tutoring charges steep hourly rates unaffordable for college students."),
        ("FRAGMENTED", "Peer coordination is scattered across WhatsApp, Discord, and Google Calendar."),
        ("ZERO SAFETY", "Informal peer swaps suffer from ghosting, no-shows, and lack of accountability.")
    ]
    sw = Inches(3.75)
    for i, (callout, exp) in enumerate(stats):
        sx = Inches(0.8) + i * Inches(3.99)
        add_card(s3, sx, Inches(1.35), sw, Inches(1.5), bg_color=LIGHT_MINT)
        tb = s3.shapes.add_textbox(sx + Inches(0.2), Inches(1.45), sw - Inches(0.4), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = callout
        p.font.name = "Cambria"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = DARK_GREEN
        p.space_after = Pt(4)
        p2 = tf.add_paragraph()
        p2.text = exp
        p2.font.name = "Calibri"
        p2.font.size = Pt(12)
        p2.font.color.rgb = DARK_TEXT

    # Bottom 4 Context Cards (2x2)
    contexts = [
        ("1. Expensive Professional Tutoring", "Commercial EdTech platforms focus on corporate professionals, charging fees beyond average student budgets."),
        ("2. Disconnected Communication Tools", "Arranging sessions through chat apps leads to missed links, scheduling clashes, and lost notes without tracking."),
        ("3. High Ghosting & No-Show Rates", "Free peer-to-peer knowledge sharing lacks formal commitment; learners or mentors cancel last-minute with zero recourse."),
        ("4. Awkward Transaction Transparency", "Students helping students hesitate to discuss hourly rates informally; lack of a clear fee ledger creates friction.")
    ]
    cw = Inches(5.7)
    ch = Inches(1.6)
    for idx, (head, body) in enumerate(contexts):
        cx = Inches(0.8) if idx % 2 == 0 else Inches(6.833)
        cy = Inches(3.15) if idx < 2 else Inches(4.95)
        add_card(s3, cx, cy, cw, ch, bg_color=CARD_BG)
        tb = s3.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.18), cw - Inches(0.4), ch - Inches(0.36))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = head
        p.font.name = "Cambria"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = DARK_GREEN
        p.space_after = Pt(4)
        p2 = tf.add_paragraph()
        p2.text = body
        p2.font.name = "Calibri"
        p2.font.size = Pt(13)
        p2.font.color.rgb = DARK_TEXT

    # ==========================================
    # SLIDE 4: Objectives & Core Modules (3 Columns)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Project Objectives & Core Modules", "Scope & Roles")
    add_footer(s4, 4)

    # Top summary card
    add_card(s4, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.95), bg_color=LIGHT_MINT)
    tb_obj = s4.shapes.add_textbox(Inches(1.0), Inches(1.45), Inches(11.333), Inches(0.75))
    tf_obj = tb_obj.text_frame
    tf_obj.word_wrap = True
    p = tf_obj.paragraphs[0]
    p.text = "Primary Goal: A trusted campus ecosystem for affordable 1-on-1 peer mentoring."
    p.font.name = "Cambria"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.space_after = Pt(2)
    p2 = tf_obj.add_paragraph()
    p2.text = "Empower students to monetize their domain expertise while enabling peers to book structured, verified learning sessions."
    p2.font.name = "Calibri"
    p2.font.size = Pt(13)
    p2.font.color.rgb = DARK_TEXT

    # 3 Role Columns below (Learner, Mentor, Admin)
    roles = [
        ("Learner Module", "For Students Seeking Guidance", [
            "Discover verified mentors by skill tags & hourly rate",
            "3-step session scheduling with open date/time chips",
            "Real-time 1-on-1 chat (5 free messages anti-spam)",
            "Track live learning hours, attended sessions & milestones"
        ]),
        ("Mentor Module", "For Student Practitioners & Tutors", [
            "Set custom hourly rates & define availability slots",
            "Showcase skills with depth ratings (Beginner to Expert)",
            "Accept, decline, or propose alternate session times",
            "Teaching Studio analytics & transparent earnings ledger"
        ]),
        ("Admin Governance", "For Platform Integrity & Moderation", [
            "Central KPI console: user counts, sessions, pending checks",
            "Member management: verify accounts, suspend bad actors",
            "Audit session lifecycle logs & manage dispute reports",
            "Financial transparency with downloadable CSV ledgers"
        ])
    ]
    rw = Inches(3.75)
    rh = Inches(4.2)
    for idx, (role_title, role_sub, pts) in enumerate(roles):
        rx = Inches(0.8) + idx * Inches(3.99)
        add_card(s4, rx, Inches(2.55), rw, rh, bg_color=CARD_BG)
        tb = s4.shapes.add_textbox(rx + Inches(0.2), Inches(2.75), rw - Inches(0.4), rh - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = role_title
        p.font.name = "Cambria"
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = DARK_GREEN
        p.space_after = Pt(2)

        p_sub = tf.add_paragraph()
        p_sub.text = role_sub
        p_sub.font.name = "Calibri"
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = MUTED_TEXT
        p_sub.space_after = Pt(14)

        for pt in pts:
            p_pt = tf.add_paragraph()
            p_pt.text = "• " + pt
            p_pt.font.name = "Calibri"
            p_pt.font.size = Pt(13)
            p_pt.font.color.rgb = DARK_TEXT
            p_pt.space_after = Pt(8)

    # ==========================================
    # SLIDE 5: Technology Stack & Database Schemas
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Technology Stack & Database Architecture", "Engineering Foundation")
    add_footer(s5, 5)

    # Top Half: 6 Tech Stack Cards (Grid of 3 x 2)
    techs = [
        ("Frontend: React 19 + Vite", "Component architecture, client routing, Lucide icons, responsive CSS variables."),
        ("Backend: Node.js + Express", "RESTful controllers, modular routes, error handling, and request validation."),
        ("Database: MongoDB Atlas", "Cloud NoSQL cluster with Mongoose ODM schemas, indexing, and references."),
        ("Real-Time: Socket.io", "Persistent bi-directional WebSocket channels for instant messaging & read status."),
        ("Auth: JWT & Bcrypt", "Stateless bearer token authorization with 10-round salted password hashing."),
        ("Ledger: Simulated Escrow", "0% platform fee, 0% tax, transparent hourly billing & manual admin settlements.")
    ]
    tw = Inches(3.75)
    th = Inches(1.3)
    for idx, (t_title, t_desc) in enumerate(techs):
        tx = Inches(0.8) + (idx % 3) * Inches(3.99)
        ty = Inches(1.35) if idx < 3 else Inches(2.8)
        add_card(s5, tx, ty, tw, th, bg_color=LIGHT_MINT)
        tb = s5.shapes.add_textbox(tx + Inches(0.15), ty + Inches(0.12), tw - Inches(0.3), th - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = t_title
        p.font.name = "Cambria"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = DARK_GREEN
        p.space_after = Pt(2)
        p2 = tf.add_paragraph()
        p2.text = t_desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.color.rgb = DARK_TEXT

    # Bottom Half: 4 Mongoose Model Cards
    models = [
        ("User Model", "name, email, password, role ('learner'|'mentor'|'admin'), bio, skillsToTeach, skillsToLearn, hourlyRate, availability, payoutUpiId, profileCompleted, isActive"),
        ("Booking Model", "learner (Ref), mentor (Ref), date, time, duration (30-120m), status ('pending'|'confirmed'|'completed'|'cancelled'), actionExpiresAt (4h), meetingUrl"),
        ("Chat Models", "Conversation: participants, bookingId, lastMessage, lastActivityAt\nMessage: conversationId, sender, text, isEdited, readBy"),
        ("Review & Wallet", "Review: mentor, learner, rating (1-5), comment\nWallet: balance (simulated top-up), user, transactions ledger")
    ]
    mw = Inches(2.8)
    mh = Inches(2.45)
    for idx, (m_title, m_fields) in enumerate(models):
        mx = Inches(0.8) + idx * Inches(2.98)
        my = Inches(4.35)
        add_card(s5, mx, my, mw, mh, bg_color=CARD_BG)
        tb = s5.shapes.add_textbox(mx + Inches(0.15), my + Inches(0.12), mw - Inches(0.3), mh - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = m_title
        p.font.name = "Cambria"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = DARK_GREEN
        p.space_after = Pt(4)
        p2 = tf.add_paragraph()
        p2.text = m_fields
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.color.rgb = DARK_TEXT
        p2.line_spacing = 1.15

    # ==========================================
    # SLIDE 6: System Architecture Diagram
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "System Architecture: End-to-End Data Flow", "High-Level Design")
    add_footer(s6, 6)

    # Reference Style: Plain transparent boxes with thick bold black borders and black arrows
    # Box 1: React Client
    b1 = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.8), Inches(3.2), Inches(4.6))
    b1.fill.background()
    b1.line.color.rgb = BLACK
    b1.line.width = Pt(2.5)
    tf1 = b1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "React Client (Frontend)\n"
    p.font.name = "Cambria"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.alignment = PP_ALIGN.CENTER
    for sp in ["Vite Development & Bundling", "React Router Navigation", "Learner & Mentor Views", "Socket.io Client Listener", "LocalStorage Token Persistence"]:
        p2 = tf1.add_paragraph()
        p2.text = "• " + sp
        p2.font.name = "Calibri"
        p2.font.size = Pt(13)
        p2.font.color.rgb = BLACK
        p2.space_after = Pt(6)

    # Arrow 1: REST
    a1 = s6.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.3), Inches(2.6), Inches(1.3), Inches(0.45))
    a1.fill.solid()
    a1.fill.fore_color.rgb = BLACK
    a1.line.color.rgb = BLACK
    tb_a1 = s6.shapes.add_textbox(Inches(4.1), Inches(2.1), Inches(1.7), Inches(0.45))
    p = tb_a1.text_frame.paragraphs[0]
    p.text = "RESTful HTTP"
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.alignment = PP_ALIGN.CENTER

    # Arrow 2: WebSockets
    a2 = s6.shapes.add_shape(MSO_SHAPE.LEFT_RIGHT_ARROW, Inches(4.3), Inches(4.8), Inches(1.3), Inches(0.45))
    a2.fill.solid()
    a2.fill.fore_color.rgb = BLACK
    a2.line.color.rgb = BLACK
    tb_a2 = s6.shapes.add_textbox(Inches(4.1), Inches(5.35), Inches(1.7), Inches(0.45))
    p = tb_a2.text_frame.paragraphs[0]
    p.text = "WebSocket (Chat)"
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.alignment = PP_ALIGN.CENTER

    # Box 2: Node.js / Express Server
    b2 = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.8), Inches(1.8), Inches(3.4), Inches(4.6))
    b2.fill.background()
    b2.line.color.rgb = BLACK
    b2.line.width = Pt(2.5)
    tf2 = b2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Node.js / Express API\n"
    p.font.name = "Cambria"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.alignment = PP_ALIGN.CENTER
    for sp in ["JWT Auth & Role Protection", "Booking Lifecycle Controller", "Socket.io Server Event Bus", "Google Meet Link Generator", "Admin Moderation Endpoints"]:
        p2 = tf2.add_paragraph()
        p2.text = "• " + sp
        p2.font.name = "Calibri"
        p2.font.size = Pt(13)
        p2.font.color.rgb = BLACK
        p2.space_after = Pt(6)

    # Arrow 3: Mongoose ODM
    a3 = s6.shapes.add_shape(MSO_SHAPE.LEFT_RIGHT_ARROW, Inches(9.4), Inches(3.7), Inches(0.9), Inches(0.45))
    a3.fill.solid()
    a3.fill.fore_color.rgb = BLACK
    a3.line.color.rgb = BLACK
    tb_a3 = s6.shapes.add_textbox(Inches(9.1), Inches(3.15), Inches(1.5), Inches(0.45))
    p = tb_a3.text_frame.paragraphs[0]
    p.text = "Mongoose ODM"
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.alignment = PP_ALIGN.CENTER

    # Box 3: MongoDB Atlas
    b3 = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.5), Inches(1.8), Inches(2.0), Inches(4.6))
    b3.fill.background()
    b3.line.color.rgb = BLACK
    b3.line.width = Pt(2.5)
    tf3 = b3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "MongoDB Atlas\n"
    p.font.name = "Cambria"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.alignment = PP_ALIGN.CENTER
    for sp in ["Users", "Bookings", "Messages", "Reviews", "Logs"]:
        p2 = tf3.add_paragraph()
        p2.text = "• " + sp
        p2.font.name = "Calibri"
        p2.font.size = Pt(13)
        p2.font.color.rgb = BLACK
        p2.space_after = Pt(8)

    # ==========================================
    # SLIDE 7: Signature Feature: 5-Step Booking Flow
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Signature Feature: 5-Stage Session Booking Lifecycle", "Core Innovation")
    add_footer(s7, 7)

    # 5 Sequential Step Cards
    steps = [
        ("Step 1", "Choose", "Learner selects duration (30-120m), date chip, and available time slot."),
        ("Step 2", "Validate", "System checks mentor schedule, prevents self-booking, sets 4h expiry."),
        ("Step 3", "Respond", "Mentor reviews request to Accept, Propose Alternate Slots, or Decline."),
        ("Step 4", "Confirm", "Session locked; Google Meet link attached; unlimited chat unlocked."),
        ("Step 5", "Complete", "Session held; earnings credited to ledger; mutual reviews submitted.")
    ]
    step_w = Inches(2.15)
    step_h = Inches(3.4)
    for idx, (st_num, st_name, st_desc) in enumerate(steps):
        st_x = Inches(0.8) + idx * Inches(2.4)
        add_card(s7, st_x, Inches(1.5), step_w, step_h, bg_color=CARD_BG)

        # Header pill inside step
        tb = s7.shapes.add_textbox(st_x + Inches(0.15), Inches(1.7), step_w - Inches(0.3), Inches(3.0))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = st_num.upper()
        p.font.name = "Calibri"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = DARK_GREEN
        p.space_after = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = st_name
        p2.font.name = "Cambria"
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = DARK_TEXT
        p2.space_after = Pt(10)

        p3 = tf.add_paragraph()
        p3.text = st_desc
        p3.font.name = "Calibri"
        p3.font.size = Pt(13)
        p3.font.color.rgb = DARK_TEXT
        p3.line_spacing = 1.2

        # Small arrow between cards (except last)
        if idx < 4:
            arr = s7.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, st_x + step_w + Inches(0.04), Inches(2.9), Inches(0.18), Inches(0.25))
            arr.fill.solid()
            arr.fill.fore_color.rgb = DARK_GREEN
            arr.line.fill.background()

    # Bottom Important Policy Callout
    add_card(s7, Inches(0.8), Inches(5.2), Inches(11.733), Inches(1.4), bg_color=LIGHT_MINT)
    tb_pol = s7.shapes.add_textbox(Inches(1.0), Inches(5.35), Inches(11.333), Inches(1.1))
    tf_pol = tb_pol.text_frame
    tf_pol.word_wrap = True
    p = tf_pol.paragraphs[0]
    p.text = "Strict 4-Hour Action Window & Counter-Proposal Protocol"
    p.font.name = "Cambria"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.space_after = Pt(4)
    p2 = tf_pol.add_paragraph()
    p2.text = "• If a mentor does not respond within 4 hours, the request automatically expires to free the learner's schedule."
    p2.font.name = "Calibri"
    p2.font.size = Pt(13)
    p2.font.color.rgb = DARK_TEXT
    p3 = tf_pol.add_paragraph()
    p3.text = "• Mentors can counter-propose up to 3 alternative time slots directly without declining the learner."
    p3.font.name = "Calibri"
    p3.font.size = Pt(13)
    p3.font.color.rgb = DARK_TEXT

    # ==========================================
    # SLIDE 8: Walkthrough: Home & Authentication
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Live Walkthrough: Landing & Authentication", "User Onboarding")
    add_footer(s8, 8)

    tb8 = s8.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.733), Inches(1.1))
    tf8 = tb8.text_frame
    tf8.word_wrap = True
    p = tf8.paragraphs[0]
    p.text = "• Dual-Path Landing: Welcomes users with 'Share what you know' (mentor) and 'Learn something new' (learner)."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT
    p.space_after = Pt(3)
    p = tf8.add_paragraph()
    p.text = "• Role-Based Registration: Simple choice between 'Learn Skills' and 'Teach & Mentor' with terms validation."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT
    p.space_after = Pt(3)
    p = tf8.add_paragraph()
    p.text = "• Secure Access: Bcrypt salted password encryption, JWT local storage, and instant redirection to role dashboards."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT

    add_image_bordered(s8, "01_home.png", Inches(0.8), Inches(2.45), width=Inches(5.7))
    add_image_bordered(s8, "02a_login.png", Inches(6.833), Inches(2.45), width=Inches(5.7))

    # ==========================================
    # SLIDE 9: Walkthrough: Discover & Profiles
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Live Walkthrough: Discovery & Profile Management", "Peer Matching")
    add_footer(s9, 9)

    tb9 = s9.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.733), Inches(1.1))
    tf9 = tb9.text_frame
    tf9.word_wrap = True
    p = tf9.paragraphs[0]
    p.text = "• Searchable Directory: Real-time search across peers and mentors by name, role, or technical skill tags."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT
    p.space_after = Pt(3)
    p = tf9.add_paragraph()
    p.text = "• Complete Profiles: Showcase bio, skills to teach, skills to learn, and weekly schedule availability."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT
    p.space_after = Pt(3)
    p = tf9.add_paragraph()
    p.text = "• Public Mentor Portfolio: Displays verified badge, hourly rates (e.g. ₹400/hr), and direct 'Book Session' action."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT

    add_image_bordered(s9, "04_discover.png", Inches(0.8), Inches(2.45), width=Inches(5.7))
    add_image_bordered(s9, "03b_other_user_profile.png", Inches(6.833), Inches(2.45), width=Inches(5.7))

    # ==========================================
    # SLIDE 10: Walkthrough: Learner & Mentor Dashboards
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Live Walkthrough: Role-Based Dashboards", "Progress & Earnings")
    add_footer(s10, 10)

    tb10 = s10.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.733), Inches(1.1))
    tf10 = tb10.text_frame
    tf10.word_wrap = True
    p = tf10.paragraphs[0]
    p.text = "• Learner Growth Tracker: Monitors live learning hours, attended sessions, mentors connected, and topic milestones."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT
    p.space_after = Pt(3)
    p = tf10.add_paragraph()
    p.text = "• Mentor Teaching Studio: Live session request flow, skill depth meters (Beginner-Expert), and student counts."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT
    p.space_after = Pt(3)
    p = tf10.add_paragraph()
    p.text = "• Transparent Earnings Ledger: Tracks lifetime revenue, pending disbursements, and inline UPI ID payout setup."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT

    add_image_bordered(s10, "05_learner_dashboard.png", Inches(0.8), Inches(2.45), width=Inches(5.7))
    add_image_bordered(s10, "06_mentor_dashboard.png", Inches(6.833), Inches(2.45), width=Inches(5.7))

    # ==========================================
    # SLIDE 11: Walkthrough: Booking Flow & Real-Time Chat
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Live Walkthrough: Session Booking & Live Chat", "Core Collaboration")
    add_footer(s11, 11)

    tb11 = s11.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.733), Inches(1.1))
    tf11 = tb11.text_frame
    tf11.word_wrap = True
    p = tf11.paragraphs[0]
    p.text = "• 3-Step Booking Wizard: Select duration (30-120 mins), choose upcoming date chip, and pick morning/evening slots."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT
    p.space_after = Pt(3)
    p = tf11.add_paragraph()
    p.text = "• Transparent Fee Model: 0% platform fee, 0% tax, transparent hourly calculation, and escrow protection notice."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT
    p.space_after = Pt(3)
    p = tf11.add_paragraph()
    p.text = "• Socket.io Duplex Chat: Real-time messaging with 5-free message anti-spam barrier until session is booked."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT

    add_image_bordered(s11, "07_booking_flow.png", Inches(0.8), Inches(2.45), width=Inches(5.7))
    add_image_bordered(s11, "08_chat.png", Inches(6.833), Inches(2.45), width=Inches(5.7))

    # ==========================================
    # SLIDE 12: Walkthrough: Admin Governance & Reviews
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Live Walkthrough: Admin Governance & Quality Control", "Governance")
    add_footer(s12, 12)

    tb12 = s12.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.733), Inches(1.1))
    tf12 = tb12.text_frame
    tf12.word_wrap = True
    p = tf12.paragraphs[0]
    p.text = "• Platform Overview KPIs: Live tracking of total members (14), mentors (8), learners (5), and sessions (29)."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT
    p.space_after = Pt(3)
    p = tf12.add_paragraph()
    p.text = "• Member Management: Administrator controls to verify student credentials, change roles, or suspend bad actors."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT
    p.space_after = Pt(3)
    p = tf12.add_paragraph()
    p.text = "• Quality & Audits: Verified post-session 1-5 star reviews, feedback collection, and downloadable financial CSV ledgers."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT

    add_image_bordered(s12, "10_admin_dashboard.png", Inches(0.8), Inches(2.45), width=Inches(5.7))
    add_image_bordered(s12, "10b_admin_members.png", Inches(6.833), Inches(2.45), width=Inches(5.7))

    # ==========================================
    # SLIDE 13: Engineering Challenge: Dark Theme Specificity
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "Engineering Challenge: Dark Theme CSS Specificity", "Technical Deep Dive")
    add_footer(s13, 13)

    # Left Column: Problem & Resolution
    add_card(s13, Inches(0.8), Inches(1.35), Inches(5.7), Inches(5.35), bg_color=CARD_BG)
    tb_ch = s13.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(5.3), Inches(4.9))
    tf_ch = tb_ch.text_frame
    tf_ch.word_wrap = True

    p = tf_ch.paragraphs[0]
    p.text = "The Problem"
    p.font.name = "Cambria"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.space_after = Pt(4)

    p = tf_ch.add_paragraph()
    p.text = "During the initial dark mode rollout, the global CSS rule inverted app root variables. While public pages rendered well, the Admin Workspace sidebar and data tables broke completely, rendering pale white backgrounds with illegible text."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT
    p.space_after = Pt(14)

    p = tf_ch.add_paragraph()
    p.text = "Root Cause Analysis"
    p.font.name = "Cambria"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.space_after = Pt(4)

    p = tf_ch.add_paragraph()
    p.text = "Unscoped `!important` declarations in `index.css` collided with admin styling variables, overriding dark slate palettes with inverted root styles."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT
    p.space_after = Pt(14)

    p = tf_ch.add_paragraph()
    p.text = "The Architectural Fix"
    p.font.name = "Cambria"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.space_after = Pt(4)

    p = tf_ch.add_paragraph()
    p.text = "Refactored administrative styling by strictly scoping them under `html[data-theme='dark'] .admin-page` and `.admin-sidebar`, guaranteeing consistent contrast without breaking public theme toggles."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT

    # Right Column: Code Box
    add_card(s13, Inches(6.833), Inches(1.35), Inches(5.7), Inches(5.35), bg_color=CODE_BG, border_color=BLACK)
    tb_code = s13.shapes.add_textbox(Inches(7.05), Inches(1.55), Inches(5.3), Inches(4.9))
    tf_code = tb_code.text_frame
    tf_code.word_wrap = True

    p = tf_code.paragraphs[0]
    p.text = "/* client/src/Pages/Admin Page/admin.css */"
    p.font.name = "Consolas"
    p.font.size = Pt(11)
    p.font.color.rgb = MUTED_TEXT
    p.space_after = Pt(8)

    code_lines = [
        ("/* Preserves intended dark emerald sidebar */", MUTED_TEXT),
        ("html[data-theme=\"dark\"] .admin-sidebar,", WHITE),
        ("html.dark .admin-sidebar {", WHITE),
        ("    background-color: #091712 !important;", CODE_GREEN),
        ("    border-right: 1px solid #1e3a2f;", WHITE),
        ("}", WHITE),
        ("", WHITE),
        ("/* Scoped workspace background & text */", MUTED_TEXT),
        ("html[data-theme=\"dark\"] .admin-page {", WHITE),
        ("    background-color: #0d1f18;", CODE_GREEN),
        ("    color: #e2f0ea;", CODE_GREEN),
        ("}", WHITE),
        ("", WHITE),
        ("html[data-theme=\"dark\"] .admin-workspace {", WHITE),
        ("    background-color: #0d1f18 !important;", CODE_GREEN),
        ("}", WHITE)
    ]
    for line, color in code_lines:
        p = tf_code.add_paragraph()
        p.text = line
        p.font.name = "Consolas"
        p.font.size = Pt(11)
        p.font.color.rgb = color
        p.space_after = Pt(2)

    # ==========================================
    # SLIDE 14: Project Conclusion & Future Scope
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "Project Conclusion & Future Roadmap", "Summary & Next Steps")
    add_footer(s14, 14)

    # Left Column: Conclusion
    add_card(s14, Inches(0.8), Inches(1.35), Inches(5.7), Inches(5.35), bg_color=CARD_BG)
    tb_con = s14.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(5.3), Inches(4.9))
    tf_con = tb_con.text_frame
    tf_con.word_wrap = True

    p = tf_con.paragraphs[0]
    p.text = "Key Project Achievements"
    p.font.name = "Cambria"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.space_after = Pt(14)

    con_bullets = [
        "Production-Grade MERN Platform: Solved peer tutoring coordination friction with modern React 19, Express REST APIs, and MongoDB Atlas.",
        "Zero-Fee Knowledge Economy: Enabled fair knowledge exchange with 0% platform fee and transparent session tracking.",
        "Deterministic Session Lifecycle: Built-in 4-hour auto-expiration and slot counter-proposals to prevent ghosting.",
        "Real-Time Collaboration: Integrated persistent Socket.io chat and automatic Google Meet video room attachment.",
        "Undergraduate Internship Success: Built collaboratively by Team CosmicCoders (Group 13) during internship at Codroid Hub Pvt. Ltd."
    ]
    for cb in con_bullets:
        p = tf_con.add_paragraph()
        p.text = "• " + cb
        p.font.name = "Calibri"
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(10)

    # Right Column: Future Scope
    add_card(s14, Inches(6.833), Inches(1.35), Inches(5.7), Inches(5.35), bg_color=LIGHT_MINT)
    tb_fut = s14.shapes.add_textbox(Inches(7.05), Inches(1.55), Inches(5.3), Inches(4.9))
    tf_fut = tb_fut.text_frame
    tf_fut.word_wrap = True

    p = tf_fut.paragraphs[0]
    p.text = "Future Technical Scope"
    p.font.name = "Cambria"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.space_after = Pt(14)

    fut_bullets = [
        ("In-Browser WebRTC Video Calling", "Build native peer-to-peer audio/video streaming directly inside the platform without external meeting providers."),
        ("Automated Razorpay Webhooks", "Implement automated bank payout webhooks and automatic refund triggers for dispute arbitration."),
        ("Push & Calendar Synchronization", "Integrate browser Service Worker push alerts and iCal calendar invites 15 minutes before sessions."),
        ("Native Mobile Application", "Develop cross-platform iOS and Android apps using React Native for instant on-the-go chat and bookings.")
    ]
    for f_title, f_desc in fut_bullets:
        p = tf_fut.add_paragraph()
        p.text = "• " + f_title + ":"
        p.font.name = "Calibri"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = DARK_GREEN
        p.space_after = Pt(2)

        p2 = tf_fut.add_paragraph()
        p2.text = f_desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(12)
        p2.font.color.rgb = DARK_TEXT
        p2.space_after = Pt(8)

    # ==========================================
    # SLIDE 15: Thank You Slide (Full Bleed Dark Green)
    # ==========================================
    s15 = prs.slides.add_slide(blank_layout)
    bg15 = s15.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg15.fill.solid()
    bg15.fill.fore_color.rgb = DARK_GREEN
    bg15.line.fill.background()

    # Logo top-right
    if os.path.exists(logo_path):
        s15.shapes.add_picture(logo_path, Inches(11.2), Inches(0.8), width=Inches(1.2))

    tb15 = s15.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(10.0), Inches(4.8))
    tf15 = tb15.text_frame
    tf15.word_wrap = True

    p = tf15.paragraphs[0]
    p.text = "Thank You!"
    p.font.name = "Cambria"
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(10)

    p = tf15.add_paragraph()
    p.text = "Questions & Discussion Welcome"
    p.font.name = "Calibri"
    p.font.size = Pt(22)
    p.font.color.rgb = RGBColor(0xDF, 0xF0, 0xE8)
    p.space_after = Pt(36)

    p = tf15.add_paragraph()
    p.text = "Project Repository: github.com/CosmicCoder06/Skill-Exchange"
    p.font.name = "Calibri"
    p.font.size = Pt(18)
    p.font.color.rgb = WHITE
    p.space_after = Pt(12)

    p = tf15.add_paragraph()
    p.text = "Live Web Deployment: https://skill-exchange-one-eta.vercel.app/"
    p.font.name = "Calibri"
    p.font.size = Pt(18)
    p.font.color.rgb = WHITE
    p.space_after = Pt(24)

    p = tf15.add_paragraph()
    p.text = "Built with ❤️ by Team CosmicCoders (Group 13)"
    p.font.name = "Calibri"
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(0xCD, 0xE5, 0xDC)

    output_path = os.path.abspath(os.path.join(base_dir, 'Skill_Exchange_Capstone_Presentation.pptx'))
    prs.save(output_path)
    print(f"Presentation generated successfully with {len(prs.slides)} slides at: {output_path}")

if __name__ == "__main__":
    build_presentation()
