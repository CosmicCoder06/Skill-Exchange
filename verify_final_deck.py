import re
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

BANNED_WORDS = [
    r"\bdeterministic\b",
    r"\becosystem\b",
    r"\becosystem-wide\b",
    r"\bbi-directional\b",
    r"\bduplex\b",
    r"\bstateless\b",
    r"\bspecificity\b",
    r"\barchitectural\b",
    r"\bgovernance\b",
    r"\bprotocol\b",
    r"\barbitration\b",
    r"\bmonetize\b",
    r"\bdisbursements\b",
    r"\bledger\b",
    r"\bfriction\b",
    r"\bempower\b",
    r"\brobust\b",
    r"\bseamless\b",
    r"\bproduction-grade\b",
    r"\bstrict\b",
    r"\bcomprehensive\b",
    r"\bleverage\b",
    r"\butilize\b",
]

EXPECTED_PICTURE_COUNTS = {
    1: 1, 2: 0, 3: 0, 4: 0, 5: 0, 6: 0, 7: 0, 8: 2, 9: 2, 10: 2, 11: 2, 12: 2, 13: 0, 14: 0, 15: 1
}

def verify():
    prs = Presentation("Skill_Exchange_Capstone_Presentation_Final.pptx")
    print(f"Total slides: {len(prs.slides)}")
    assert len(prs.slides) == 15, f"Expected 15 slides, got {len(prs.slides)}"

    # 1. Picture counts
    for idx, slide in enumerate(prs.slides):
        s_num = idx + 1
        pic_count = sum(1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
        expected = EXPECTED_PICTURE_COUNTS.get(s_num, 0)
        assert pic_count == expected, f"Slide {s_num} picture count {pic_count} != expected {expected}"
        print(f"Slide {s_num}: Pictures = {pic_count} (MATCH)")

    # 2. Slide 1 empty fields
    s1_text = "\n".join(p.text for s in prs.slides[0].shapes if s.has_text_frame for p in s.text_frame.paragraphs)
    assert "Team Leader:   ___" in s1_text, "Team Leader empty field missing"
    assert "Team Members: ___" in s1_text, "Team Members empty field missing"
    assert "Project Guide: ___" in s1_text, "Project Guide empty field missing"
    print("Slide 1 fields: All empty with underscores (MATCH)")

    # 3. Slide 15 heart emoji check
    s15_text = "\n".join(p.text for s in prs.slides[14].shapes if s.has_text_frame for p in s.text_frame.paragraphs)
    assert "❤️" not in s15_text and "\u2764" not in s15_text, "Heart emoji still present on slide 15!"
    assert "Built by Team CosmicCoders (Group 13)" in s15_text, "Slide 15 attribution text mismatch"
    print("Slide 15 heart emoji: Removed successfully (MATCH)")

    # 4. Banned words check across all slides
    banned_found = []
    for idx, slide in enumerate(prs.slides):
        s_num = idx + 1
        all_text = " ".join(p.text for s in slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs)
        for pattern in BANNED_WORDS:
            matches = re.findall(pattern, all_text, re.IGNORECASE)
            if matches:
                banned_found.append((s_num, pattern, matches))

    if banned_found:
        print("FAIL: Banned words found:")
        for s_num, pattern, matches in banned_found:
            print(f"  Slide {s_num}: {pattern} -> {matches}")
        assert False, "Banned words detected!"
    else:
        print("Banned words check: 0 banned words found across all 15 slides! (MATCH)")

    # 5. Speaker notes check
    for s_idx in range(1, 14): # Slides 2 to 14 (indices 1 to 13)
        slide = prs.slides[s_idx]
        note_text = slide.notes_slide.notes_text_frame.text.strip()
        sentences = [s.strip() for s in re.split(r'[.!?]', note_text) if s.strip()]
        assert len(sentences) >= 3, f"Slide {s_idx+1} speaker note has only {len(sentences)} sentences: '{note_text}'"
        print(f"Slide {s_idx+1} Speaker Note: {len(sentences)} sentences (MATCH)")

    # 6. Specific synopsis content checks
    s3_text = " ".join(p.text for s in prs.slides[2].shapes if s.has_text_frame for p in s.text_frame.paragraphs)
    assert "NO CENTRAL PLATFORM" in s3_text, "Slide 3 missing NO CENTRAL PLATFORM"
    assert "HARD TO FIND MENTORS" in s3_text, "Slide 3 missing HARD TO FIND MENTORS"
    assert "NO ORGANIZED SYSTEM" in s3_text, "Slide 3 missing NO ORGANIZED SYSTEM"
    print("Slide 3 synopsis check: (MATCH)")

    s4_text = " ".join(p.text for s in prs.slides[3].shapes if s.has_text_frame for p in s.text_frame.paragraphs)
    assert "Learner Module" in s4_text and "Mentor Module" in s4_text and "Admin" in s4_text, "Slide 4 roles mismatch"
    print("Slide 4 synopsis check: (MATCH)")

    s5_text = " ".join(p.text for s in prs.slides[4].shapes if s.has_text_frame for p in s.text_frame.paragraphs)
    assert "REST APIs" in s5_text, "Slide 5 missing REST APIs"
    assert "Escrow" not in s5_text, "Slide 5 still mentions Escrow"
    print("Slide 5 synopsis check: (MATCH)")

    s14_text = " ".join(p.text for s in prs.slides[13].shapes if s.has_text_frame for p in s.text_frame.paragraphs)
    assert "AI-Based Mentor Recommendations" in s14_text, "Slide 14 missing AI recommendations"
    assert "Video Call Integration" in s14_text, "Slide 14 missing video call integration"
    assert "Skill Certification Badges" in s14_text, "Slide 14 missing certification badges"
    assert "Gamified Learning Paths" in s14_text, "Slide 14 missing gamified learning paths"
    print("Slide 14 synopsis check: (MATCH)")

    print("\nALL FINAL VERIFICATION CHECKS PASSED PERFECTLY!")

if __name__ == "__main__":
    verify()
