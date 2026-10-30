from pptx import Presentation

prs = Presentation('Skill_Exchange_Capstone_Presentation.pptx')

for s_idx, slide in enumerate(prs.slides):
    for sh_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                if len(p.runs) > 1:
                    runs_desc = [f'R{r_idx}("{r.text}", bold={r.font.bold}, size={r.font.size})' for r_idx, r in enumerate(p.runs)]
                    print(f'Slide {s_idx+1} Shape {sh_idx} P{p_idx}: {" | ".join(runs_desc)}')
